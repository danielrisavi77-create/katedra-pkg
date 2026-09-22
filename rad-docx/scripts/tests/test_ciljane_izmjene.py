#!/usr/bin/env python3
"""Synthetic regressions: editable target, contextual degrees and scoped OOXML edits.

No student files, real signatures or external image/transcription services are used.
"""
from __future__ import annotations

import copy
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import etree
from PIL import Image

SCRIPT = Path(__file__).resolve().parents[1] / 'ciljane_izmjene.py'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def rewrite_part(path, name, fn):
    data = parts(path)
    data[name] = fn(data[name])
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as z:
        for n, b in data.items():
            z.writestr(n, b)


def fixture(path):
    """Two explicit cover blocks + cards/CV/body. Everything is invented."""
    d = Document()
    d.add_paragraph('IZMIŠLJENO SVEUČILIŠTE')
    d.add_paragraph('Primjer Autora')
    d.add_paragraph('POKAZNI RAD O UČENJU')
    d.add_page_break()
    d.add_paragraph('IZMIŠLJENO SVEUČILIŠTE')
    d.add_paragraph('Primjer Autora')
    d.add_paragraph('POKAZNI RAD O UČENJU')
    d.add_paragraph('Mentorica: Primjer Mentorice')
    d.add_page_break()
    d.add_paragraph('TEMELJNA DOKUMENTACIJSKA KARTICA')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'Akademski naziv'
    p = t.cell(0, 1).paragraphs[0]
    p.add_run('magistra').bold = True
    p.add_run(' pokazne discipline').italic = True
    t.cell(1, 0).text = 'Kratica akademskoga naziva'
    t.cell(1, 1).text = 'mag. test.'
    d.add_paragraph('BASIC DOCUMENTATION CARD')
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = 'Academic Degree'
    t.cell(0, 1).text = 'Master of Demonstration'
    d.add_paragraph('10. ŽIVOTOPIS')
    p = d.add_paragraph()
    p.add_run('2020.–2023. ').bold = True
    p.add_run('sveučilišna prvostupnica pokazne discipline').italic = True
    d.add_paragraph('SAŽETAK')
    d.add_paragraph('Uzorak ima 203 sudionika; nalaz 54,0 % (1).')
    d.add_section(WD_SECTION_START.NEW_PAGE)
    d.add_paragraph('1. UVOD')
    d.add_paragraph('Tijelo rada, uz sačuvane tablice, citate i priloge.')
    d.sections[-1].footer.is_linked_to_previous = False
    p = d.sections[-1].footer.paragraphs[0]
    f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = '1'
    r.append(text); f.append(r); p._p.append(f)
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:start'), '1')
    d.sections[-1]._sectPr.append(pg)
    buf = io.BytesIO()
    Image.new('RGB', (64, 16), (70, 70, 70)).save(buf, format='PNG')
    buf.seek(0)
    d.add_paragraph('SINTETIČKI PRILOG (nije potpis)')
    d.add_picture(buf, width=Inches(1))
    d.save(path)


class TargetedEdits(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if not SCRIPT.exists():
            raise AssertionError('Targeted OOXML entry point is not implemented')
        spec = importlib.util.spec_from_file_location('ciljane_izmjene', SCRIPT)
        cls.m = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.m)

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.src = self.root / 'source.docx'
        self.out = self.root / 'edited.docx'
        self.snap = self.root / 'snapshot.docx'
        fixture(self.src)

    def plan(self, *ops):
        return {'version': 1, 'target_sha256': sha(self.src),
                'output_kind': 'docx', 'operations': list(ops)}

    def index(self, exact):
        ps = self.m.inspect_document(self.src)['paragraphs']
        return next(p['index'] for p in ps if p['text'] == exact)

    def replace(self, old, new, paragraph=None):
        return {'op': 'replace_text', 'paragraph': self.index(paragraph or old),
                'old': old, 'new': new}

    def degree(self, context, old='magistra', new='sveučilišna magistra',
               paragraph='magistra pokazne discipline'):
        return {**self.replace(old, new, paragraph), 'op': 'replace_degree',
                'context': context,
                'source': {'kind': 'user_confirmed', 'locator': 'synthetic decision 1'}}

    def remove(self):
        return {'op': 'remove_first_cover', 'end_block': 3,
                'expected_text': ['IZMIŠLJENO SVEUČILIŠTE', 'Primjer Autora',
                                  'POKAZNI RAD O UČENJU']}

    def apply(self, plan):
        return self.m.apply_plan(self.src, self.out, plan, snapshot=self.snap)

    def test_editable_output_and_original_identity(self):
        p = self.plan(self.replace('Mentorica: Primjer Mentorice',
                                   'Mentorica: Primjer Mentorice, dr. sc.'))
        p['references'] = [{'role': 'visual_reference', 'id': 'different-cover-photo'}]
        r = self.apply(p)
        d = Document(self.out)
        self.assertIn('Primjer Autora', '\n'.join(x.text for x in d.paragraphs))
        self.assertNotIn('different-cover-photo', parts(self.out)['word/document.xml'].decode())
        self.assertEqual(r['status'], 'structural_pass')
        self.assertEqual(r['before_sha256'], sha(self.src))
        self.assertEqual(r['after_sha256'], sha(self.out))
        self.assertIsNone(r['visual_check']['ok'])

    def test_generated_image_cannot_be_output(self):
        p = self.plan(self.remove()); p['output_kind'] = 'image'
        with self.assertRaises(self.m.EditError): self.apply(p)
        self.assertFalse(self.out.exists())

    def test_image_cannot_be_target(self):
        self.src.write_bytes(b'not a Word document')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_renaming_extension_does_not_create_word(self):
        bad = self.root / 'output.png'
        with self.assertRaises(self.m.EditError):
            self.m.apply_plan(self.src, bad, self.plan(self.remove()), snapshot=self.snap)
        self.assertFalse(bad.exists())

    def test_awarded_degree_keeps_earned_cv_and_styles(self):
        old = parts(self.src)
        self.apply(self.plan(self.degree('awarded')))
        root = etree.fromstring(parts(self.out)['word/document.xml'])
        srcroot = etree.fromstring(old['word/document.xml'])
        self.assertEqual(root.xpath('//w:rPr', namespaces=NS).__len__(),
                         srcroot.xpath('//w:rPr', namespaces=NS).__len__())
        self.assertEqual([etree.tostring(x) for x in root.xpath('//w:rPr', namespaces=NS)],
                         [etree.tostring(x) for x in srcroot.xpath('//w:rPr', namespaces=NS)])
        self.assertIn('sveučilišna prvostupnica pokazne discipline',
                      parts(self.out)['word/document.xml'].decode())
        self.assertIn('sveučilišna magistra', parts(self.out)['word/document.xml'].decode())
        for n in old:
            if n != 'word/document.xml': self.assertEqual(old[n], parts(self.out)[n], n)

    def test_earned_degree_edit_does_not_touch_card(self):
        op = self.degree('earned', 'sveučilišna prvostupnica pokazne discipline',
                         'sveučilišna prvostupnica pokazne discipline (univ. bacc. test.)',
                         '2020.–2023. sveučilišna prvostupnica pokazne discipline')
        self.apply(self.plan(op))
        self.assertEqual(Document(self.out).tables[0].cell(0, 1).text,
                         'magistra pokazne discipline')

    def test_degree_context_mismatch_rejected(self):
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.degree('earned')))
        self.assertFalse(self.out.exists())

    def test_degree_without_provenance_rejected(self):
        op = self.degree('awarded'); del op['source']
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_unconfirmed_audio_not_degree_evidence(self):
        op = self.degree('awarded'); op['source']['kind'] = 'unconfirmed_audio'
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_english_awarded_card_supported_without_translation(self):
        op = self.degree('awarded', 'Master of Demonstration', 'Master of Example',
                         'Master of Demonstration')
        self.apply(self.plan(op))
        self.assertEqual(Document(self.out).tables[1].cell(0, 1).text, 'Master of Example')

    def test_degree_card_other_row_rejected(self):
        op = self.degree('awarded', 'Akademski naziv', 'Drugo', 'Akademski naziv')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_remove_first_cover_preserves_exact_remaining_xml(self):
        before = parts(self.src)
        root = etree.fromstring(before['word/document.xml'])
        body = root.find('w:body', NS)
        expected = [etree.tostring(x) for x in list(body)[4:]]
        self.apply(self.plan(self.remove()))
        after = parts(self.out)
        outbody = etree.fromstring(after['word/document.xml']).find('w:body', NS)
        self.assertEqual(expected, [etree.tostring(x) for x in outbody])
        for n in before:
            if n != 'word/document.xml': self.assertEqual(before[n], after[n], n)
        self.assertEqual(self.src.read_bytes(), self.snap.read_bytes())

    def test_two_covers_cannot_be_removed_in_one_operation(self):
        op = self.remove(); op['end_block'] = 8
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_wrong_expected_cover_rejected(self):
        op = self.remove(); op['expected_text'][1] = 'Different person'
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_removal_without_explicit_boundary_rejected(self):
        d = Document(self.src)
        for p in d.paragraphs:
            for b in p._p.findall('.//w:br', NS): b.getparent().remove(b)
        d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_section_break_in_removed_cover_rejected(self):
        d = Document(self.src)
        d.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
        d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_bookmark_on_removed_cover_rejected(self):
        d = Document(self.src)
        b = OxmlElement('w:bookmarkStart'); b.set(qn('w:id'), '42'); b.set(qn('w:name'), 'x')
        d.paragraphs[0]._p.append(b); d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_first_page_header_policy_rejected(self):
        d = Document(self.src); d.sections[0].different_first_page_header_footer = True
        d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_duplicate_remove_operation_rejected(self):
        with self.assertRaises(self.m.EditError):
            self.apply(self.plan(self.remove(), self.remove()))

    def test_replace_and_remove_bind_to_original_indices(self):
        self.apply(self.plan(self.remove(), self.degree('awarded')))
        self.assertEqual(Document(self.out).tables[0].cell(0, 1).text,
                         'sveučilišna magistra pokazne discipline')

    def test_stale_source_hash_rejected(self):
        p = self.plan(self.remove())
        d = Document(self.src); d.add_paragraph('Later edit'); d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(p)
        self.assertFalse(self.snap.exists()); self.assertFalse(self.out.exists())

    def test_existing_output_not_overwritten(self):
        self.out.write_bytes(b'KEEP')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))
        self.assertEqual(self.out.read_bytes(), b'KEEP')

    def test_snapshot_cannot_be_original_or_output(self):
        for snap in (self.src, self.out):
            with self.subTest(snap=snap):
                with self.assertRaises(self.m.EditError):
                    self.m.apply_plan(self.src, self.out, self.plan(self.remove()), snapshot=snap)
        self.assertFalse(self.out.exists())

    def test_unknown_operation_and_fields_fail_closed(self):
        for op in ({'op': 'remove_all_covers'}, {**self.remove(), 'allow_everything': True}):
            with self.subTest(op=op):
                with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_ambiguous_text_rejected(self):
        d = Document(self.src); d.add_paragraph('word word'); d.save(self.src)
        with self.assertRaises(self.m.EditError):
            self.apply(self.plan(self.replace('word', 'new', 'word word')))

    def test_split_runs_with_identical_format_preserved(self):
        d = Document(self.src); p = d.add_paragraph()
        p.add_run('univ. ba').bold = True; p.add_run('cc. test.').bold = True; d.save(self.src)
        self.apply(self.plan(self.replace('bacc.', 'mag.', 'univ. bacc. test.')))
        p = Document(self.out).paragraphs[-1]
        self.assertEqual(p.text, 'univ. mag. test.')
        self.assertEqual([r.bold for r in p.runs], [True, True])

    def test_mixed_style_match_rejected(self):
        with self.assertRaises(self.m.EditError):
            self.apply(self.plan(self.replace('magistra pokazne', 'magistra druge',
                                              'magistra pokazne discipline')))

    def test_fields_are_not_plain_text_targets(self):
        d = Document(self.src); p = d.add_paragraph('OLD')
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'DATE'); p._p.append(f)
        d.save(self.src)
        with self.assertRaises(self.m.EditError):
            self.apply(self.plan(self.replace('OLD', 'NEW')))

    def test_track_changes_never_autoaccepted(self):
        d = Document(self.src)
        ins = OxmlElement('w:ins'); d.paragraphs[-2]._p.append(ins); d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_valid_scope_verifies(self):
        p = self.plan(self.degree('awarded'))
        self.apply(p)
        self.assertTrue(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_outside_scope_numeric_change_detected(self):
        p = self.plan(self.degree('awarded')); self.apply(p)
        rewrite_part(self.out, 'word/document.xml', lambda b: b.replace(b'203', b'204'))
        self.assertFalse(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_media_change_detected(self):
        p = self.plan(self.remove()); self.apply(p)
        media = next(n for n in parts(self.out) if n.startswith('word/media/'))
        rewrite_part(self.out, media, lambda b: b + b'CHANGED')
        self.assertFalse(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_removed_second_cover_detected(self):
        p = self.plan(self.remove()); self.apply(p)
        def corrupt(b):
            root = etree.fromstring(b); body = root.find('w:body', NS)
            for node in list(body)[:5]: body.remove(node)
            return etree.tostring(root)
        rewrite_part(self.out, 'word/document.xml', corrupt)
        self.assertFalse(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_formatting_drift_detected(self):
        p = self.plan(self.degree('awarded')); self.apply(p)
        rewrite_part(self.out, 'word/document.xml', lambda b: b.replace(b'<w:i/>', b'<w:i w:val="0"/>'))
        self.assertFalse(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_cli_report_and_inspection(self):
        planfile = self.root / 'plan.json'; report = self.root / 'report.json'
        planfile.write_text(json.dumps(self.plan(self.remove())), encoding='utf-8')
        r = subprocess.run([sys.executable, str(SCRIPT), 'apply', str(self.src),
                            '--plan', str(planfile), '--out', str(self.out),
                            '--snapshot', str(self.snap), '--report', str(report)],
                           capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertEqual(json.loads(report.read_text(encoding='utf-8'))['after_sha256'], sha(self.out))
        r = subprocess.run([sys.executable, str(SCRIPT), 'inspect', str(self.out)],
                           capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 0, r.stderr)
        self.assertEqual(json.loads(r.stdout)['target_sha256'], sha(self.out))

    def test_cli_duplicate_json_keys_rejected(self):
        planfile = self.root / 'plan.json'
        planfile.write_text('{"version":1,"version":2}', encoding='utf-8')
        r = subprocess.run([sys.executable, str(SCRIPT), 'apply', str(self.src),
                            '--plan', str(planfile), '--out', str(self.out),
                            '--snapshot', str(self.snap), '--report', str(self.root / 'r.json')],
                           capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 2)
        self.assertFalse(self.out.exists())


    def test_plain_numbered_heading_ends_cv_context(self):
        d = Document(self.src)
        anchor = next(p for p in d.paragraphs if p.text == 'SAŽETAK')
        h = OxmlElement('w:p'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = '11. PRILOZI'; r.append(t); h.append(r); anchor._p.addprevious(h)
        p = OxmlElement('w:p'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = 'prvostupnica u citatu'; r.append(t); p.append(r); anchor._p.addprevious(p)
        d.save(self.src)
        op = self.degree('earned', 'prvostupnica', 'magistra', 'prvostupnica u citatu')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_signed_package_not_silently_invalidated(self):
        with zipfile.ZipFile(self.src, 'a') as z:
            z.writestr('_xmlsignatures/sig1.xml', '<signature/>')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))
        self.assertFalse(self.snap.exists())

    def test_report_failure_removes_only_new_output_not_snapshot(self):
        real = self.m._write_new
        report = self.root / 'report.json'
        def fail_report(path, content):
            if Path(path) == report: raise PermissionError('simulated report failure')
            return real(path, content)
        with mock.patch.object(self.m, '_write_new', side_effect=fail_report):
            with self.assertRaises(self.m.EditError):
                self.m.apply_plan(self.src, self.out, self.plan(self.remove()),
                                  snapshot=self.snap, report=report)
        self.assertTrue(self.snap.exists())
        self.assertFalse(self.out.exists())

    def test_output_is_verified_after_write(self):
        real = self.m._write_new
        def corrupt_after_write(path, content):
            result = real(path, content)
            if Path(path) == self.out:
                rewrite_part(path, 'word/document.xml', lambda b: b.replace(b'203', b'204'))
            return result
        with mock.patch.object(self.m, '_write_new', side_effect=corrupt_after_write):
            with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))
        self.assertFalse(self.out.exists())

    def test_invalid_xml_character_is_actionable_error(self):
        op = self.replace('magistra', 'magistra\x00', 'magistra pokazne discipline')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))
        self.assertFalse(self.out.exists())

    def test_boolean_paragraph_index_not_integer(self):
        op = self.replace('Primjer Autora', 'Drugi autor'); op['paragraph'] = True
        with self.assertRaises(self.m.EditError): self.apply(self.plan(op))

    def test_metadata_change_is_outside_scope(self):
        p = self.plan(self.remove()); self.apply(p)
        rewrite_part(self.out, 'docProps/core.xml', lambda b: b.replace(b'</cp:coreProperties>',
                     b'<cp:category>changed</cp:category></cp:coreProperties>'))
        self.assertFalse(self.m.verify_plan(self.src, self.out, p)['ok'])

    def test_duplicate_zip_entries_rejected(self):
        import warnings
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with zipfile.ZipFile(self.src, 'a') as z:
                z.writestr('word/document.xml', b'<invalid/>')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_dangling_output_symlink_not_followed(self):
        self.out.symlink_to(self.root / 'missing.docx')
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))
        self.assertTrue(self.out.is_symlink())

    def test_cv_year_entry_does_not_end_cv_region(self):
        d = Document(self.src)
        anchor = next(p for p in d.paragraphs if p.text == 'SAŽETAK')
        p = OxmlElement('w:p'); r = OxmlElement('w:r'); t = OxmlElement('w:t')
        t.text = 'Stečeni naziv: prvostupnica'; r.append(t); p.append(r)
        anchor._p.addprevious(p); d.save(self.src)
        op = self.degree('earned', 'prvostupnica', 'sveučilišna prvostupnica',
                         'Stečeni naziv: prvostupnica')
        self.apply(self.plan(op))
        self.assertIn('Stečeni naziv: sveučilišna prvostupnica',
                      parts(self.out)['word/document.xml'].decode())

    def test_legacy_paragraph_assignment_reproduces_run_loss(self):
        legacy = Document(self.src)
        p = legacy.tables[0].cell(0, 1).paragraphs[0]
        self.assertEqual(len(p._p.findall('w:r/w:rPr', NS)), 2)
        p.text = p.text.replace('magistra', 'sveučilišna magistra')
        self.assertEqual(len(p._p.findall('w:r/w:rPr', NS)), 0)
        self.apply(self.plan(self.degree('awarded')))
        fixed = Document(self.out).tables[0].cell(0, 1).paragraphs[0]
        self.assertEqual(len(fixed._p.findall('w:r/w:rPr', NS)), 2)

    def test_table_cell_revision_requires_explicit_view(self):
        d = Document(self.src)
        d.tables[0].cell(0, 0)._tc.get_or_add_tcPr().append(OxmlElement('w:cellDel'))
        d.save(self.src)
        with self.assertRaises(self.m.EditError): self.apply(self.plan(self.remove()))

    def test_reference_id_unicode_does_not_leave_unverified_output(self):
        p = self.plan(self.remove())
        p['references'] = [{'role': 'visual_reference', 'id': 'id-\ud800'}]
        self.assertEqual(self.apply(p)['status'], 'structural_pass')

    def test_xml_comment_does_not_break_inventory(self):
        def insert_comment(b):
            root = etree.fromstring(b)
            root.find('w:body', NS).append(etree.Comment('test comment'))
            return etree.tostring(root)
        rewrite_part(self.src, 'word/document.xml', insert_comment)
        self.assertTrue(self.m.inspect_document(self.src)['paragraphs'])


if __name__ == '__main__':
    unittest.main()
