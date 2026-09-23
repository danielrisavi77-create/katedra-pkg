#!/usr/bin/env python3
"""Synthetic delivery regression cases; no private manuscript or signed data."""
import copy
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import warnings
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from PIL import Image, PngImagePlugin

SCRIPTS = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SCRIPTS))
SPEC = importlib.util.find_spec('delivery_integrity')
if SPEC:
    import delivery_integrity as di


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def rewrite(path, replacements=None, additions=None):
    with zipfile.ZipFile(path) as archive:
        data = {x: archive.read(x) for x in archive.namelist()}
    data.update(additions or {})
    for key, func in (replacements or {}).items():
        data[key] = func(data[key])
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        for key, raw in data.items():
            archive.writestr(key, raw)


class DeliveryTests(unittest.TestCase):
    def setUp(self):
        self.assertIsNotNone(SPEC, 'D delivery integrity module is required')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.before = self.root / 'before.docx'
        self.after = self.root / 'after.docx'
        doc = Document()
        doc.add_heading('Rezultati', 1)
        doc.add_paragraph('U analiziranim slučajevima ukupno je 600 t (Autor, 2020).')
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Godina'
        table.cell(0, 1).text = 'Otpad (t)'
        table.cell(1, 0).text = '2020'
        table.cell(1, 1).text = '600'
        doc.add_paragraph('Izvor: sintetički primjer.')
        doc.save(self.before)
        self.after.write_bytes(self.before.read_bytes())

    def comparison(self, **kw):
        return di.compare_documents(self.before, self.after, **kw)

    def test_identical_document_has_scoped_structural_pass(self):
        r = self.comparison()
        self.assertEqual(r['status'], 'pass')
        self.assertFalse(r['whole_document'])
        self.assertEqual(r['visual_status'], 'unmeasured')

    def test_typographic_formatting_preserves_ordered_content(self):
        doc = Document(self.after)
        doc.paragraphs[1].runs[0].font.size = Pt(12)
        doc.paragraphs[1].runs[0].bold = True
        doc.tables[0].style = 'Light Shading Accent 1'
        doc.save(self.after)
        self.assertEqual(self.comparison()['status'], 'pass')

    def test_run_splitting_without_text_change_is_not_semantic_mutation(self):
        doc = Document(self.after)
        p = doc.paragraphs[1]
        text = p.text
        p.clear()
        p.add_run(text[:12]).bold = True
        p.add_run(text[12:])
        doc.save(self.after)
        self.assertEqual(self.comparison()['status'], 'pass')

    def test_number_citation_caption_and_table_order_changes_block(self):
        changes = [lambda d: setattr(d.paragraphs[1], 'text', 'Ukupno je 750 t (Autor, 2020).'),
                   lambda d: setattr(d.paragraphs[-1], 'text', 'Izvor: drugi autor.'),
                   lambda d: setattr(d.tables[0].cell(1, 1), 'text', '601')]
        for change in changes:
            with self.subTest(change=change):
                doc = Document(self.before)
                change(doc)
                doc.save(self.after)
                self.assertEqual(self.comparison()['status'], 'blocked')

    def test_field_flattening_is_not_formatting(self):
        doc = Document(self.before)
        p = doc.add_paragraph()
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'REF Example')
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = '600'
        r.append(t); f.append(r); p._p.append(f)
        doc.save(self.before)
        doc = Document(self.before); doc.paragraphs[-1].text = '600'; doc.save(self.after)
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_hyperlink_target_change_blocks(self):
        self._hyperlink('https://example.org/one')
        self.after.write_bytes(self.before.read_bytes())
        rewrite(self.after, {'word/_rels/document.xml.rels': lambda b: b.replace(b'/one', b'/two')})
        self.assertEqual(self.comparison()['status'], 'blocked')

    def _hyperlink(self, target):
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        doc = Document(self.before)
        rel = doc.part.relate_to(target, RT.HYPERLINK, is_external=True)
        link = OxmlElement('w:hyperlink'); link.set(qn('r:id'), rel)
        run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'Izvor'; run.append(t); link.append(run)
        doc.add_paragraph()._p.append(link); doc.save(self.before)

    def pictures(self):
        doc = Document(self.before)
        for index, rgb in enumerate(((20, 60, 120), (210, 80, 30)), 1):
            path = self.root / f'{index}.png'
            Image.new('RGB', (100, 60), rgb).save(path)
            shape = doc.add_picture(str(path), width=Inches(2))
            shape._inline.docPr.set('descr', f'Sintetički prikaz {index}')
        doc.save(self.before); self.after.write_bytes(self.before.read_bytes())

    def test_visual_alt_text_bound_to_identity_not_position(self):
        self.pictures()
        doc = Document(self.after)
        a, b = [x._inline.docPr for x in doc.inline_shapes]
        a.set('descr', 'Sintetički prikaz 2'); b.set('descr', 'Sintetički prikaz 1')
        doc.save(self.after)
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_visual_resize_allowed_but_changed_pixels_block(self):
        self.pictures()
        doc = Document(self.after); doc.inline_shapes[0].width = Inches(2.5); doc.save(self.after)
        self.assertEqual(self.comparison()['status'], 'pass')
        buf = io.BytesIO(); Image.new('RGB', (100, 60), (1, 2, 3)).save(buf, format='PNG')
        rewrite(self.after, additions={'word/media/image1.png': buf.getvalue()})
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_changed_chart_data_blocks_even_when_prose_unchanged(self):
        chart = b'<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"><c:chart><c:v>600</c:v></c:chart></c:chartSpace>'
        rewrite(self.before, additions={'word/charts/chart1.xml': chart})
        self.after.write_bytes(self.before.read_bytes())
        rewrite(self.after, {'word/charts/chart1.xml': lambda b: b.replace(b'600', b'601')})
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_section_width_and_image_width_overflow_detected(self):
        self.pictures()
        doc = Document(self.after)
        doc.sections[0].left_margin = Inches(3); doc.sections[0].right_margin = Inches(3)
        doc.inline_shapes[0].width = Inches(5); doc.save(self.after)
        codes = {x['code'] for x in di.geometry_findings(self.after)}
        self.assertIn('table_width', codes); self.assertIn('image_width', codes)

    def test_later_landscape_section_not_judged_by_first_section(self):
        from docx.enum.section import WD_SECTION
        doc = Document(self.before)
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.page_width = Inches(14); section.page_height = Inches(8.5)
        section.left_margin = Inches(0.5); section.right_margin = Inches(0.5)
        table = doc.add_table(rows=1, cols=2); table.autofit = False
        table.columns[0].width = Inches(5); table.columns[1].width = Inches(5)
        doc.save(self.after)
        self.assertNotIn('table_width', {x['code'] for x in di.geometry_findings(self.after)})

    def test_fixed_row_and_long_token_are_render_risks_not_guaranteed_pass(self):
        doc = Document(self.after); row = doc.tables[0].rows[1]
        pr = row._tr.get_or_add_trPr(); height = OxmlElement('w:trHeight')
        height.set(qn('w:val'), '220'); height.set(qn('w:hRule'), 'exact'); pr.append(height)
        row.cells[0].text = 'https://example.org/' + 'abcdefghij' * 15
        doc.save(self.after)
        codes = {x['code'] for x in di.geometry_findings(self.after)}
        self.assertIn('fixed_row_height', codes); self.assertIn('unbroken_token', codes)

    def test_named_metadata_cleanup_preserves_content_and_snapshot(self):
        doc = Document(self.before); doc.core_properties.author = 'Synthetic Author'; doc.save(self.before)
        result = di.clean_metadata(self.before, self.root/'clean.docx', self.root/'snapshot.docx', ['core', 'application', 'thumbnail'])
        self.assertEqual(result['status'], 'pass')
        self.assertEqual(self.before.read_bytes(), (self.root/'snapshot.docx').read_bytes())
        self.assertEqual(di.compare_documents(self.before, self.root/'clean.docx', mode='metadata_only')['status'], 'pass')
        self.assertFalse(di.metadata_findings(self.root/'clean.docx', ['core', 'application', 'thumbnail']))
        self.assertIn('embedded_metadata', result['unchecked'])

    def test_metadata_reintroduced_by_save_does_not_pass_privacy(self):
        result = di.clean_metadata(self.before, self.root/'clean.docx', self.root/'snapshot.docx', ['core'])
        doc = Document(self.root/'clean.docx'); doc.core_properties.author = 'New Editor'; doc.save(self.after)
        self.assertTrue(di.metadata_findings(self.after, ['core']))

    def test_metadata_cleanup_never_overwrites_existing_output(self):
        old = self.after.read_bytes()
        with self.assertRaises(ValueError):
            di.clean_metadata(self.before, self.after, self.root/'snapshot.docx', ['core'])
        self.assertEqual(old, self.after.read_bytes())

    def test_unknown_metadata_group_rejected_instead_of_claiming_all(self):
        with self.assertRaises(ValueError):
            di.clean_metadata(self.before, self.root/'clean.docx', self.root/'snapshot.docx', ['all'])
        self.assertFalse((self.root/'clean.docx').exists())

    def test_duplicate_zip_parts_rejected(self):
        with warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning)
            with zipfile.ZipFile(self.after, 'a') as archive:
                archive.writestr('docProps/core.xml', b'<x/>')
        with self.assertRaises(ValueError): self.comparison()

    def test_revision_and_comments_never_implicitly_accepted(self):
        doc = Document(self.before)
        ins = OxmlElement('w:ins'); ins.set(qn('w:author'), 'Synthetic Editor')
        r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = 'new'; r.append(t); ins.append(r)
        doc.paragraphs[1]._p.append(ins); doc.save(self.before)
        with self.assertRaises(ValueError):
            di.clean_metadata(self.before, self.root/'clean.docx', self.root/'snapshot.docx', ['core'])
        self.assertFalse((self.root/'clean.docx').exists())

    def test_dangling_relationship_is_invalid(self):
        rewrite(self.after, {'word/_rels/document.xml.rels': lambda b: b.replace(b'Target="styles.xml"', b'Target="missing.xml"')})
        with self.assertRaises(ValueError): self.comparison()

    def test_metadata_bound_content_control_prevents_silent_deletion(self):
        doc = Document(self.before)
        f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'DOCPROPERTY Author')
        doc.paragraphs[1]._p.append(f); doc.save(self.before)
        with self.assertRaises(ValueError):
            di.clean_metadata(self.before, self.root/'clean.docx', self.root/'snapshot.docx', ['core'])

    def test_metadata_only_mode_rejects_font_and_geometry_changes(self):
        doc = Document(self.after); doc.paragraphs[1].runs[0].bold = True; doc.save(self.after)
        self.assertEqual(self.comparison(mode='metadata_only')['status'], 'blocked')

    def test_text_inside_drawing_is_protected(self):
        self.pictures()
        from lxml import etree
        ns = 'http://schemas.openxmlformats.org/drawingml/2006/main'
        doc = Document(self.before)
        text = etree.Element('{'+ns+'}t'); text.text = '600 t'
        doc.inline_shapes[0]._inline.append(text); doc.save(self.before)
        self.after.write_bytes(self.before.read_bytes())
        rewrite(self.after, {'word/document.xml': lambda b: b.replace(b'600 t</a:t>', b'750 t</a:t>')})
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_unclassified_embedded_content_change_is_protected(self):
        rewrite(self.before, additions={'word/afchunk.html': b'<p>600 t</p>'})
        self.after.write_bytes(self.before.read_bytes())
        rewrite(self.after, additions={'word/afchunk.html': b'<p>750 t</p>'})
        self.assertEqual(self.comparison()['status'], 'blocked')

    def test_missing_pdf_measurement_does_not_allow_visual_pass(self):
        p, _ = self.render_fixture(); approval = self.attest(p)
        with patch.object(di, '_pdf_page_count', side_effect=ValueError('pdfinfo unavailable'), create=True):
            self.assertEqual(di.verify_render(self.after, p, approval)['status'], 'unmeasured')

    def test_pdf_actual_page_count_must_match_image_manifest(self):
        p, _ = self.render_fixture(); approval = self.attest(p)
        with patch.object(di, '_pdf_page_count', return_value=2, create=True):
            self.assertNotEqual(di.verify_render(self.after, p, approval)['status'], 'pass')

    def render_fixture(self):
        # Synthetic evidence tests contract integrity; this is NOT an actual renderer QA.
        counter = patch.object(di, '_pdf_page_count', return_value=1, create=True)
        counter.start(); self.addCleanup(counter.stop)
        directory = self.root/'render'; directory.mkdir()
        pdf = directory/'document.pdf'; pdf.write_bytes(b'%PDF-synthetic-contract-fixture')
        image = directory/'page-1.png'; Image.new('RGB', (1200, 1600), 'white').save(image)
        manifest = {'schema_version':1, 'kind':'katedra_render', 'document_sha256':sha(self.after),
                    'renderer':{'name':'Synthetic test renderer', 'version':'fixture-only'},
                    'pdf':{'path':'document.pdf', 'sha256':sha(pdf)},
                    'pages':[{'index':1,'path':'page-1.png','sha256':sha(image),'width':1200,'height':1600}],
                    'created_at':'2026-09-23T07:00:00+00:00'}
        path = directory/'render.json'; path.write_text(json.dumps(manifest), encoding='utf-8')
        return path, manifest

    def test_missing_render_never_means_visual_pass(self):
        self.assertEqual(di.verify_render(self.after, None, None)['status'], 'unmeasured')

    def test_pending_visual_review_is_unmeasured(self):
        p, _ = self.render_fixture()
        self.assertEqual(di.verify_render(self.after, p, None)['status'], 'unmeasured')

    def attest(self, manifest):
        approval = {'schema_version':1,'kind':'katedra_visual_review','render_manifest_sha256':sha(manifest),
                    'document_sha256':sha(self.after),'reviewer':'Synthetic QA',
                    'reviewed_pages':[1],'decision':'approved','notes':'Synthetic contract only, not real renderer.',
                    'reviewed_at':'2026-09-23T07:01:00+00:00'}
        p = self.root/'render'/'review.json'; p.write_text(json.dumps(approval),encoding='utf-8')
        return p

    def test_bound_named_all_page_review_passes_only_declared_scope(self):
        p, _ = self.render_fixture(); approval = self.attest(p)
        r = di.verify_render(self.after,p,approval)
        self.assertEqual(r['status'],'pass')
        self.assertEqual(r['renderer']['name'],'Synthetic test renderer')
        self.assertFalse(r['independent_verification'])

    def test_stale_document_pdf_page_and_manifest_are_rejected(self):
        p, data = self.render_fixture(); approval = self.attest(p)
        for target in (self.after, p.parent/'document.pdf', p.parent/'page-1.png', p):
            with self.subTest(target=target.name):
                original = target.read_bytes(); target.write_bytes(original+b' ')
                self.assertNotEqual(di.verify_render(self.after,p,approval)['status'],'pass')
                target.write_bytes(original)

    def test_partial_page_review_cannot_pass(self):
        p, data = self.render_fixture(); approval = self.attest(p)
        data['pages'].append({**data['pages'][0], 'index':2})
        p.write_text(json.dumps(data),encoding='utf-8')
        a = json.loads(approval.read_text()); a['render_manifest_sha256']=sha(p)
        approval.write_text(json.dumps(a),encoding='utf-8')
        self.assertNotEqual(di.verify_render(self.after,p,approval)['status'],'pass')

    def test_render_manifest_cannot_reference_files_outside_its_directory(self):
        p, data = self.render_fixture(); data['pdf']['path']='../before.docx'
        data['pdf']['sha256']=sha(self.before); p.write_text(json.dumps(data),encoding='utf-8')
        approval=self.attest(p)
        self.assertNotEqual(di.verify_render(self.after,p,approval)['status'],'pass')

    def test_delivery_without_render_not_released_even_after_content_parity(self):
        manifest = {'schema_version':1,'kind':'katedra_delivery','before':{'path':'before.docx','sha256':sha(self.before)},
                    'after':{'path':'after.docx','sha256':sha(self.after)},'mode':'format_only',
                    'metadata_groups':[], 'render_manifest':None,'visual_review':None,
                    'excluded':[{'item':'JMBAG','reason':'Author will supply before submission'}]}
        p=self.root/'delivery.json'; p.write_text(json.dumps(manifest),encoding='utf-8')
        r=di.check_delivery(p)
        self.assertEqual(r['status'],'unmeasured')
        self.assertEqual(r['excluded'],manifest['excluded'])
        self.assertFalse(r['whole_document'])

if __name__ == '__main__': unittest.main()
