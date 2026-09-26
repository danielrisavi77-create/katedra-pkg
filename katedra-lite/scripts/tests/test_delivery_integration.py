#!/usr/bin/env python3
"""D CLI and gate integration; render availability is never presumed."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import gate
ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT/'rad-docx'/'scripts'))
import delivery_integrity as di
from docx import Document


class DeliveryIntegrationTests(unittest.TestCase):
    def test_phase_gate_opt_in_and_exact_target_binding(self):
        base={'rad':'rad.docx','pdf':None,'profil':'p.json','tip':'zavrsni','kat':'.katedra'}
        for phase in ('audit','predaja'):
            self.assertNotIn('delivery_integrity',[x.kid for x in gate.koraci(phase,base)])
            steps=gate.koraci(phase,{**base,'delivery_manifest':'.katedra/delivery.json','project_root':str(Path.cwd())})
            selected=[x for x in steps if x.kid=='delivery_integrity']
            self.assertEqual(len(selected),1)
            self.assertTrue(selected[0].blokira)
            self.assertEqual(selected[0].satelit,'rad-docx')
            self.assertIn('--rad',selected[0].argv)
            self.assertIn('rad.docx',selected[0].treba)

    def test_cli_writes_unmeasured_not_pass_for_missing_render(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);state=root/'.katedra';state.mkdir()
            before=root/'before.docx';after=root/'after.docx'
            doc=Document();doc.add_paragraph('Sintetički rad, 600 t.');doc.save(before)
            after.write_bytes(before.read_bytes())
            manifest={'schema_version':1,'kind':'katedra_delivery','before':{'path':'before.docx','sha256':di._sha(before)},
                      'after':{'path':'after.docx','sha256':di._sha(after)},'mode':'format_only','metadata_groups':[],
                      'render_manifest':None,'visual_review':None,'excluded':[]}
            m=state/'delivery.json';m.write_text(json.dumps(manifest),encoding='utf-8')
            out=state/'result.json'
            cmd=[sys.executable,str(Path(di.__file__)),'check','--manifest',str(m),'--project-root',str(root),
                 '--rad',str(after),'--out',str(out)]
            p=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',timeout=30)
            self.assertEqual(p.returncode,2,p.stdout+p.stderr)
            self.assertTrue(out.is_file(),p.stdout+p.stderr)
            self.assertEqual(json.loads(out.read_text(encoding='utf-8'))['status'],'unmeasured')
            old=out.read_bytes()
            p=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',timeout=30)
            self.assertEqual(p.returncode,2)
            self.assertEqual(out.read_bytes(),old)

    def test_unavailable_renderer_is_not_a_release_pass(self):
        from unittest.mock import patch
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);document=root/'doc.docx';Document().save(document)
            with patch.object(di.shutil,'which',return_value=None):
                result=di.render_document(document,root/'render')
            self.assertEqual(result['status'],'unmeasured')
            self.assertFalse((root/'render').exists())

if __name__=='__main__':unittest.main()
