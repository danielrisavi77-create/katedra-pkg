#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Ograde za kvarove sa sesije 22. 9. 2026. (završni rad, Veleučilište Baltazar
Zaprešić, APA; dorada po 12 primjedbi mentorice poslanih mailom).

Svaki test gleda ponašanje, ne tekst datoteke (v. test_stari_kvarovi.py, kvar 141):
  K1  stanje_init: --fakultet-izvan-registryja bez --fakultet ruši pun init
  K2  extract_comments: primjedbe iz maila nisu imale ulaz u zamjerke.json
  K3  check_ai_style: gomilanje ograda („…, a ne …", „ne dokazuje") nije se mjerilo
  K4  revizije.py toc: Wordov sadržaj u w:sdt tiho preskočen; JMBAG čitan kao redak
  K5  verify_rewrite: namjerni ispravak brojke blokirao je jednako kao tihi
"""
import json
import os
import subprocess
import sys
import tempfile

TU = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(TU)
KORIJEN = os.path.dirname(SCRIPTS)
FIX = os.path.join(KORIJEN, "tests", "fixtures", "ucenje_2026_09_22")
sys.path.insert(0, SCRIPTS)

PALO, SVE = [], []


def check(naziv, uvjet, detalj=""):
    SVE.append(naziv)
    print("  %-8s %s" % ("✓" if uvjet else "✗ FAIL", naziv))
    if not uvjet:
        PALO.append(naziv)
        if detalj:
            print("           detalj: %r" % (detalj,))


def pokreni(args, cwd):
    r = subprocess.run([sys.executable] + args, cwd=cwd, capture_output=True,
                       text=True, encoding="utf-8")
    return r.returncode, r.stdout + r.stderr


def k1_stanje_init():
    with tempfile.TemporaryDirectory() as tmp:
        kod, izlaz = pokreni([os.path.join(SCRIPTS, "stanje_init.py"), "--mod", "audit",
                              "--tip", "zavrsni", "--tema", "Test",
                              "--fakultet-izvan-registryja", "bak-zapresic",
                              "--ogranicenje", "nema profila", "--ima", "rad_docx"], tmp)
        check("K1: izvan-registryja bez --fakultet prolazi pun init", kod == 0, izlaz[-300:])
        stanje = os.path.join(tmp, ".katedra", "stanje.json")
        slug = ""
        if os.path.isfile(stanje):
            with open(stanje, encoding="utf-8") as f:
                slug = (json.load(f).get("fakultet") or {}).get("slug", "")
        check("K1: slug upisan iz --fakultet-izvan-registryja", slug == "bak-zapresic", slug)
        kod2, _ = pokreni([os.path.join(SCRIPTS, "stanje_init.py"), "--mod", "audit",
                           "--tip", "zavrsni", "--tema", "Test",
                           "--fakultet-izvan-registryja", "bak-zapresic", "--force"], tmp)
        check("K1: ograničenje je i dalje obavezno", kod2 == 2)


def k2_mail():
    with tempfile.TemporaryDirectory() as tmp:
        out = os.path.join(tmp, "zamjerke.json")
        kod, izlaz = pokreni([os.path.join(SCRIPTS, "extract_comments.py"),
                              os.path.join(FIX, "mail_mentorice.txt"),
                              "--autor", "Ana Anić", "--out", out], tmp)
        d = {}
        if os.path.isfile(out):
            with open(out, encoding="utf-8") as f:
                d = json.load(f)
        zs = d.get("zamjerke", [])
        check("K2: 12 primjedbi iz maila → 12 zamjerki", len(zs) == 12, (kod, izlaz[-300:]))
        if len(zs) == 12:
            check("K2: primjedba ne povlači citirani mail ispod sebe",
                  "U redu, hvala" not in zs[11]["tekst"], zs[11]["tekst"])
            check("K2: datum „22. ruj” nije stavka",
                  all(not z["tekst"].startswith("ruj") for z in zs))
            check("K2: autor i izvor_id upisani",
                  zs[0]["autor"] == "Ana Anić"
                  and zs[0]["izvor_id"].startswith("tekst:mail_mentorice.txt:1"))


def k3_ograde():
    import check_ai_style as S
    ograden = ("Autori nalaze povezanost, a ne uzročnost. To ne dokazuje učinak. "
               "Riječ je o procjeni, a ne o mjerenju. Podatak nije dokaz utjecaja. "
               "Nalaz ne jamči ishod, a ne dokazuje ni smjer. ") * 3
    cist = ("Autori nalaze pozitivnu povezanost u presječnom nacrtu. Učinak je mjeren "
            "samoprocjenom, pa se tumači oprezno. Model objašnjava dio varijance. ") * 6
    a1 = S.analiza([ograden])
    a2 = S.analiza([cist])
    n1 = [p for _, p, _ in S.nalazi(a1) if p.startswith("ograde")]
    n2 = [p for _, p, _ in S.nalazi(a2) if p.startswith("ograde")]
    check("K3: gusto ograđen tekst dobiva nalaz", bool(n1), a1.get("ograde_na_1000"))
    check("K3: oprezan tekst bez gomilanja ograda nema nalaz", not n2, a2.get("ograde_na_1000"))


def _docx_sa_sdt_sadrzajem(put):
    """Najmanji .docx u kojem je sadržaj u w:sdt (kao što ga Word sprema) i
    naslovnica s retkom „kolegij<TAB>JMBAG"."""
    from docx import Document
    from docx.oxml import parse_xml
    d = Document()
    d.add_paragraph("Bihevioralni aspekti poslovanja\t0123456789")
    W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    redci = "".join(
        f'<w:p><w:pPr><w:pStyle w:val="TOC{1 if i < 2 else 2}"/></w:pPr>'
        f'<w:hyperlink w:anchor="_Toc{i}"><w:r><w:t>{n}</w:t><w:tab/><w:t>{b}</w:t></w:r>'
        f'</w:hyperlink></w:p>'
        for i, (n, b) in enumerate([("SAŽETAK", 1), ("1. UVOD", 3), ("1.1. Predmet", 3)]))
    sdt = parse_xml(f'<w:sdt {W}><w:sdtPr><w:docPartObj><w:docPartGallery w:val="Table of Contents"/>'
                    f'</w:docPartObj></w:sdtPr><w:sdtContent>{redci}</w:sdtContent></w:sdt>')
    d.element.body.insert(1, sdt)
    d.save(put)


def k4_toc():
    import docx
    import revizije as R
    with tempfile.TemporaryDirectory() as tmp:
        put = os.path.join(tmp, "toc.docx")
        _docx_sa_sdt_sadrzajem(put)
        e = R._find_toc_paragraphs(docx.Document(put))
        naslovi = [x[1] for x in e]
        check("K4: redci sadržaja unutar w:sdt se vide", naslovi == ["SAŽETAK", "1. UVOD", "1.1. Predmet"],
              naslovi)
        check("K4: „kolegij<TAB>JMBAG” nije redak sadržaja",
              not any("Bihevioralni" in n for n in naslovi), naslovi)


def _docx(put, redci):
    from docx import Document
    d = Document()
    for r in redci:
        d.add_paragraph(r)
    d.save(put)


def k5_namjerno():
    with tempfile.TemporaryDirectory() as tmp:
        a, b = os.path.join(tmp, "a.docx"), os.path.join(tmp, "b.docx")
        _docx(a, ["Za angažiranost potpora je bila slabija, na razini p < 0,10 (Levitats, 2019)."])
        _docx(b, ["Regresijski model pokazao je značajan učinak (p = 0,02) (Levitats, 2019)."])
        vr = os.path.join(SCRIPTS, "verify_rewrite.py")
        kod0, izl0 = pokreni([vr, "--zahvat", "stil", a, b], tmp)
        check("K5: nenavedena izmjena brojke i dalje blokira", kod0 == 1, izl0[-200:])
        kod1, izl1 = pokreni([vr, "--zahvat", "stil",
                              "--namjerno", "0,10=ispravak prema izvorniku, str. 845",
                              "--namjerno", "0,02=ispravak prema izvorniku, str. 845", a, b], tmp)
        check("K5: navedena namjerna izmjena ne blokira brojke",
              "brojke odstupaju" not in izl1 and "namjerne izmjene brojki: 2" in izl1, izl1[-400:])
        kod2, izl2 = pokreni([vr, "--zahvat", "stil", "--namjerno", "0,02", a, b], tmp)
        check("K5: namjerno bez razloga se odbija", kod2 != 0 and "RAZLOG" in izl2, izl2[-200:])


def main():
    print("UČENJE 22. 9. 2026. — Baltazar, mentorica mailom")
    k1_stanje_init()
    k2_mail()
    k3_ograde()
    k4_toc()
    k5_namjerno()
    print(f"\n{len(SVE) - len(PALO)}/{len(SVE)} prošlo")
    return 1 if PALO else 0


if __name__ == "__main__":
    sys.exit(main())
