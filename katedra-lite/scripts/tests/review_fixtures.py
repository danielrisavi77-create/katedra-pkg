import json
from pathlib import Path
from docx import Document
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from artifact_state import record_artifact
from claim_ledger import zapisi_jsonl
from evidence_model import file_sha256, stable_source_id, stable_evidence_id, stable_claim_id, text_sha256

def make_project(root: Path):
    state=root/'.katedra'; state.mkdir(parents=True,exist_ok=True)
    sources=root/'izvori'; sources.mkdir(exist_ok=True)
    text='Otpad iznosi 20 t.'; source=sources/'synthetic.txt'; source.write_text(text+'\n',encoding='utf-8')
    sid=stable_source_id({'autor':'Synthetic','godina':2020,'naslov':'Testni izvještaj'})
    locator={'kind':'passage','passage':'odlomak 1'}; eid=stable_evidence_id(sid,locator,text)
    ev={'schema_version':1,'source_id':sid,'evidence_id':eid,'text':text,'text_sha256':text_sha256(text),'locator':locator,'source_path':'izvori/synthetic.txt','source_sha256':file_sha256(source)}
    claims=[]; doc=Document()
    for chapter in ('4','6'):
        location={'chapter':chapter,'paragraph':'1'}
        claims.append({'schema_version':1,'claim_id':stable_claim_id(text,location),'text':text,'location':location,'evidence':[{'evidence_id':eid,'relation':'supports'}]})
        doc.add_heading(chapter+'. Testno poglavlje',level=1); doc.add_paragraph(text)
    document=root/'rad.docx'; doc.save(document)
    zapisi_jsonl(state/'claims.jsonl',claims); zapisi_jsonl(state/'evidence.jsonl',[ev])
    snapshot={'izvori':[{'source_id':sid,'verification':{'status':'verified'},'quality':{'class':'primary'}}]}
    (state/'izvori.json').write_text(json.dumps(snapshot),encoding='utf-8')
    record_artifact(root,document)
    return document,state
