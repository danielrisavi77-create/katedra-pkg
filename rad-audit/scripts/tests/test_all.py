#!/usr/bin/env python3
"""Regresijski testovi za rad-audit skripte (nova/izmijenjena logika).

Uporaba:  python3 test_all.py

Gradi fixture (make_fixtures.py) u privremeni folder, poziva funkcije skripti
IZRAVNO (bez subprocessa) hvatajući stdout, i provjerava očekivane ishode.
Ne zahtijeva pytest — čisti stdlib + python-docx (isto ograničenje kao i
ostatak skilla). Exit kod 0 = sve prošlo, 1 = barem jedan test pao.
"""
import sys
import os
import io
import contextlib
import tempfile
import shutil

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

import make_fixtures  # noqa: E402

RESULTS = []


def check(name, condition, detail=""):
    RESULTS.append((name, bool(condition), detail))


def capture(fn, *args):
    buf = io.StringIO()
    code = 0
    try:
        with contextlib.redirect_stdout(buf):
            code = fn(*args) or 0
    except SystemExit as e:
        code = e.code if isinstance(e.code, int) else 0
    return buf.getvalue(), code



def test_r16_vancouver():
    """R16 — Vancouver (N). Dijalekt je bio DOKUMENTIRAN, a nije postojao u kodu:
    common.py nije imao ni riječ 'vancouver', check_citations.py je primao samo
    jedan argument, testova nije bilo, a manifest je tvrdio sposobnost
    hr.citations.vancouver.v1. Na stvarnom radu (HKS-FZS) to je davalo
    '0 citata' + 1 lažni kritični nalaz. Ovi testovi postoje da se to ne ponovi."""
    from common import detect_citation_style, find_vancouver_citations, LIT_HEADING_RE

    t = "Skrb poboljšava kvalitetu (1). Stigma je opisana (12,40) i drugdje (5)."
    stil, brojaci = detect_citation_style(t)
    check("R16: (N) tekst se detektira kao vancouver", stil == "vancouver", f"{stil} {brojaci}")
    check("R16: [N] tekst se i dalje detektira kao ieee",
          detect_citation_style("Nosivost je provjerena [1] i [12].")[0] == "ieee")

    n = lambda x, **k: [b for _p, br in find_vancouver_citations(x, **k) for b in br]
    check("R16: svezak(broj) 53(3-4) NIJE citat", n("Služba Božja. 2013;53(3-4):367-76.") == [])
    check("R16: godina (2003) NIJE citat", n("Recommendation Rec(2003)24 Vijeća Europe") == [])
    check("R16: n (%) u tablici NIJE citat", n("158 (77,8)", u_tablici=True) == [])
    check("R16: citat iza broja U PROZI jest citat", n("korelacija iznosi 0,53 (21).") == [21])
    check("R16: grupa (12,40) daje oba broja", n("stigma (12,40)") == [12, 40])

    check("R16: 'POPIS CITIRANE LITERATURE' je prepoznat naslov",
          bool(LIT_HEADING_RE.search("8. POPIS CITIRANE LITERATURE")))
    check("R16: 'Izvori i literatura' je prepoznat naslov",
          bool(LIT_HEADING_RE.search("Izvori i literatura")))



def test_r14_r15():
    """R14/R15 su bili opisani u SKILL.md bez ijednog testa (nađeno mehanički,
    zakrpa.py --provjeri-tvrdnje). Ovi testovi zatvaraju tu rupu."""
    import check_citations_authoryear as CAY
    check("R14: 'Izvori i literatura' je prepoznat naslov popisa",
          bool(CAY.HEADING_RE.search("Izvori i literatura")))
    check("R14: 'POPIS LITERATURE' i dalje prepoznat",
          bool(CAY.HEADING_RE.search("POPIS LITERATURE")))
    check("R14: padež stranog prezimena (Lipskom ~ Lipsky)",
          CAY.isti_kljuc("lipskom", "lipsky") if hasattr(CAY, "isti_kljuc")
          else "lipsky" in CAY._osnova("lipskom") or "lipsk" in CAY._osnova("lipsky"),
          (CAY._osnova("lipskom"), CAY._osnova("lipsky")))

    from apply_safe_fixes import fix_quotes_by_paragraph
    xml = ('<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
           '<w:r><w:t>Program „Neovisno življenje" i pojam "drugi navod" u istom odlomku.'
           '</w:t></w:r></w:p>')
    out = fix_quotes_by_paragraph(xml)
    izlaz = out[0] if isinstance(out, tuple) else out
    check("R15: toggle vidi već otvoreni „ (6/6, ne 11/1)",
          izlaz.count("„") == 2 and izlaz.count("”") == 2, izlaz)



def test_v195():
    """Regresije uvedene 5.9.2026. (kvarovi 61–71).

    Svaki test ovdje odgovara jednom kvaru koji je propuštao pogreške u predanu
    verziju. Pravilo: kvar bez izvršivog testa vraća se.
    """
    import common as C
    import check_citations_authoryear as AY
    import numbers_inventory as NI
    import check_typography as TY
    import check_placeholders as PH

    # Kvar 67 — prefiksno stapanje ključeva
    siroces, bez_ref = AY.uskladi_kljuceve(
        {("markov", "2021")}, {("marković", "2021")})
    check("R17: Markov i Marković NISU isti autor",
          ("marković", "2021") in siroces and ("markov", "2021") in bez_ref,
          (siroces, bez_ref))
    siroces, bez_ref = AY.uskladi_kljuceve(
        {("markovića", "2021")}, {("marković", "2021")})
    check("R17: padež (Markovića ~ Marković) i dalje spaja",
          not siroces and not bez_ref, (siroces, bez_ref))

    # Kvar 68 — narativni citat ne prelazi odlomak
    kljucevi = C.parse_ay_narrative("1. UVOD\nPrema Beckeru (2007) tržište reagira.")
    check("R18: naslov iznad citata ne postaje autor",
          ("uvod", "2007") not in kljucevi, kljucevi)

    # Kvar 70 — uvodna riječ nije prezime
    check("R18: 'Prema Kovačević (2019)' daje ključ kovačević",
          ("kovačević", "2019") in C.parse_ay_narrative(
              "Prema Kovačević (2019) izlaznost raste."),
          C.parse_ay_narrative("Prema Kovačević (2019) izlaznost raste."))

    # Kvar 64 — cijeli broj uz imenicu
    n = NI.uzorak_nalazi(["Obuhvaćeno je 147 ispitanika.",
                          "U analizu je ušlo 152 ispitanika."])
    check("R19: 147 vs 152 ispitanika je nalaz", len(n) == 1, n)
    n = NI.uzorak_nalazi(["Obuhvaćeno je 147 ispitanika.",
                          "Svih 147 ispitanika je odgovorilo."])
    check("R19: ista brojka dvaput NIJE nalaz", n == [], n)

    # Kvar 65 — zbroj kategorija je pozvan iz main(), ne mrtav kod
    check("R19: zbroj_kategorija dostupan iz modula",
          callable(getattr(NI, "zbroj_kategorija", None)))
    check("R19: uzorak_nalazi dostupan iz modula",
          callable(getattr(NI, "uzorak_nalazi", None)))

    # Kvar 66 — tipografija koju referenca traži
    for ime, izvor, trag in [
        ("duga crtica", "Tekst \u2014 umetak \u2014 dalje.", "duga crtica"),
        ("trotočka", "Tekst se nastavlja...", "tri točke"),
        ("nedosljedan postotak", "Udio 45% i udio 62 % rada.", "postotak nedosljedno"),
    ]:
        nalazi = _tipografija_nalazi(TY, izvor)
        check(f"R20: {ime} je nalaz",
              any(trag in x for x in nalazi), nalazi)

    # Kvar 71 — radne oznake, uključujući fusnote
    check("R21: check_placeholders postoji i ima uzorke",
          bool(PH.UZORCI) and callable(PH.nalazi))

    # Kvar 69 — jedan LIT_HEADING_RE za sve tri točke
    import check_overlap as CO
    import check_repetition as CR  # noqa: F401
    check("R21: check_overlap koristi zajednički LIT_HEADING_RE",
          CO.LIT_HEADING_RE is C.LIT_HEADING_RE)



def test_stvarni_rad():
    """Kvarovi 75–78 — lažni nalazi nađeni na stvarnom radu (FPZG, 5 172 riječi).

    Prvi puni prolaz novoga gatea kroz stvarni rad dao je 2 kritična i 2
    kozmetička nalaza, a nijedan nije bio greška u radu. Gate koji tako puca
    prestaje se čitati, pa je svaki od njih zaslužio test.
    """
    import common as C
    import check_citations_authoryear as AY
    import check_typography as TY

    # 75 — posvojni pridjev od prezimena
    for oblik, prezime in [("putnamovo", "putnam"), ("putnamovim", "putnam"),
                           ("closina", "closa"), ("thinusinu", "thinus")]:
        siroces, bez_ref = AY.uskladi_kljuceve({(oblik, "2021")}, {(prezime, "2021")})
        check(f"R23: posvojni pridjev {oblik} ~ {prezime}",
              not siroces and not bez_ref, (oblik, prezime, siroces, bez_ref))

    # ne smije spojiti dva različita autora
    siroces, bez_ref = AY.uskladi_kljuceve({("putnamovo", "2021")}, {("kovac", "2021")})
    check("R23: posvojni pridjev NE spaja različita prezimena",
          bool(siroces and bez_ref), (siroces, bez_ref))

    # 76 — institucionalni autor s malom riječi u sredini
    aliasi = AY.biblio_aliasi("Notes from Poland (2026) President vetoes bill.")
    check("R24: 'Notes from Poland' daje alias poland → notes",
          aliasi.get(("poland", "2026")) == ("notes", "2026"), aliasi)
    kljucevi, _ = AY.extract_biblio_keys("Notes from Poland (2026) President vetoes bill.")
    check("R24: alias NE stvara drugi unos u popisu",
          len(kljucevi) == 1, kljucevi)
    # regresija zakrpe za 76: sirovi ključ se dodaje samo za institucionalna
    # imena, nikad kad je prva riječ uvodna („Prema Marković" ≠ ključ `prema`)
    k = C.parse_ay_narrative("Prema Marković (2021) izlaznost raste.")
    check("R24: 'Prema Marković' ne daje ključ 'prema'",
          ("prema", "2021") not in k and ("marković", "2021") in k, k)
    # Narativni uzorak iz „Notes from Poland" uhvati samo „Poland"; vezu na
    # jedinicu iz popisa uspostavlja alias, ne parser. Test ide end-to-end.
    k = C.parse_ay_narrative("Izvor: Notes from Poland (2026).")
    aliasi = AY.biblio_aliasi("Notes from Poland (2026) President vetoes bill.")
    razrijeseni = {aliasi.get(x, x) for x in k}
    check("R24: alias vezuje 'poland' iz teksta na 'notes' iz popisa",
          ("notes", "2026") in razrijeseni, (k, razrijeseni))

    # 77 — apostrof u engleskom naslovu nije polunavodnik
    n = _tipografija_nalazi(TY, "Alcidi (2026) Orb\u00e1n\u2019s defeat opens the door to EU funds.")
    check("R25: apostrof (U+2019 među slovima) NIJE nalaz",
          not any("polunavodnic" in x for x in n), n)
    n = _tipografija_nalazi(TY, "Sanctions in the \u2018Rule of Law\u2019 Conflict.")
    check("R25: pravi polunavodnici JESU nalaz",
          any("polunavodnic" in x for x in n), n)

    # 87 — crta u umetnutom položaju je ISPRAVNA u hrvatskom
    n = _tipografija_nalazi(TY, "Rad \u2013 uz ogradu \u2013 pokazuje pomak.")
    check("R30: crta u umetnutom polo\u017eaju NIJE nalaz (hrvatski)",
          not any("umetnut" in x for x in n), n)
    n = _tipografija_nalazi(TY, "Rad \u2014 uz ogradu \u2014 pokazuje pomak.")
    check("R30: duga crtica JEST nalaz", any("duga crtica" in x for x in n), n)

    # 90 — engleska posvojna množina nije polunavodnik
    n = _tipografija_nalazi(TY, "Scientists\u2019 warning on affluence. University students\u2019 behavior.")
    check("R30: posvojna mno\u017eina (students\u2019) NIJE polunavodnik",
          not any("polunavodnic" in x for x in n), n)

    # 88 — veličina podskupine nije proturječje ukupnog uzorka
    import numbers_inventory as NI2
    n = NI2.uzorak_nalazi([
        "U analizu je u\u0161lo 172 ispitanika.",
        "Prema spolu, 114 ispitanika bile su \u017eene.",
        "Odgovor je dalo 45 ispitanika u dobi do 21 godine."])
    check("R31: veli\u010dine podskupina NISU proturje\u010dje", n == [], n)
    n = NI2.uzorak_nalazi([
        "U analizu je u\u0161lo 172 ispitanika.",
        "Ukupno je sudjelovalo 180 ispitanika."])
    check("R31: dvije razli\u010dite veli\u010dine UKUPNOG uzorka JESU nalaz", len(n) == 1, n)

    # 89 — sitni brojevi se zbrajaju slučajno
    check("R31: 'ukupno 3 = 1+1+0+1' nije skup kategorija",
          NI2.zbroj_kategorija(["Bilo je ukupno 3 slu\u010daja: 1 i 1 te 0 i 1."]) == [],
          NI2.zbroj_kategorija(["Bilo je ukupno 3 slu\u010daja: 1 i 1 te 0 i 1."]))

    # 91 — rad koji citira u fusnotama ne prolazi kroz autor-godina siročad
    fus = ("Vidi \u010dl. 141. Ustava, NN 56/90. Ibid., str. 214. "
           "Barbi\u0107, J., Pravo dru\u0161tava, Zagreb, 2013., str. 87. "
           "Op. cit. (bilj. 4), str. 91. Nav. dj. (bilj. 6), t. 130.")
    tijelo = "Na\u010delo vladavine prava sadr\u017eano je u \u010dl. 3. Ustava."
    check("R32: fusnotni aparat + tijelo bez oznaka = fusnotno citiranje",
          C.detect_footnote_citing(fus, tijelo), (fus[:60], tijelo))
    check("R32: rad s autor-godina citatima NIJE fusnotni",
          not C.detect_footnote_citing(
              fus, "Prema Marković (2021) i Horvat (2018) te (Kos, 2023) nalaz stoji."),
          "autor-godina u tijelu")
    check("R32: bez fusnota nije fusnotni",
          not C.detect_footnote_citing("", tijelo))
    check("R32: fusnota bez citatnog aparata nije fusnotni",
          not C.detect_footnote_citing(
              "Autorica zahvaljuje mentoru na strpljenju tijekom izrade rada.", tijelo))

    # 92 — raspodjela postotaka nije sukob
    import io as _io, contextlib as _cx
    from unittest import mock as _mock
    import numbers_inventory as NI3

    def _inv(tekst):
        buf = _io.StringIO()
        with _mock.patch.object(NI3, "load_docx_text",
                                lambda p, include_tables=True: (tekst, [], None)):
            with _cx.redirect_stdout(buf):
                NI3.main("(test)")
        return buf.getvalue()

    raspodjela = " ".join(
        f"Samo {v} % ispitanika dalo je taj odgovor." for v in
        ["12,6", "13,8", "15,9", "20,1", "23,5", "28,1", "29,1", "32,0"])
    izlaz = _inv(raspodjela)
    check("R33: osam postotaka uz 'samo' NIJE sukob",
          "'samo' + %" not in izlaz, izlaz[-400:])
    check("R33: raspodjela se izrijekom broji, ne prešućuje",
          "raspodjela, ne sukob" in izlaz or "nema" in izlaz, izlaz[-300:])

    sukob = ("Udio zaposlenih iznosio je 12,5 %. "
             "U drugom poglavlju udio zaposlenih iznosi 17,5 %.")
    izlaz = _inv(sukob)
    check("R33: dvije vrijednosti uz sadr\u017eajni pojam JESU sukob",
          "zaposlenih" in izlaz and "⚠" in izlaz, izlaz[-400:])

    # 93–97 — oblici popisa literature nađeni na stvarnim tehničkim radovima
    JEDINICE = [
        ("93: godina na kraju retka",
         "Hrvatski zavod za javno zdravstvo. Hrvatski dan mo\u017edanog udara. "
         "Zagreb: HZJZ, 2024.", ("hrvatski", "2024")),
        ("93: institucija s malim rije\u010dima u nazivu",
         "Dr\u017eavni zavod za statistiku. Istra\u017eivanje i razvoj u 2023., "
         "priop\u0107enje ZTI-2024-2-1. Zagreb: DZS, 2024.", ("dr\u017eavni", "2024")),
        ("97: godina uz broj sveska",
         "Bezak, B., Kova\u010di\u0107, S. i Brali\u0107, M. Mehani\u010dka trombektomija. "
         "Medicina Fluminensis. Vol. 57, br. 4(2021), str. 328\u2013340.", ("bezak", "2021")),
        ("97: strani \u010dasopis s to\u010dkom iza broja",
         "Luengo-Fernandez, R. i Leal, J. Economic burden of stroke across Europe. "
         "European Stroke Journal. Vol. 5, br. 1(2020), str. 17\u201325.",
         ("luengo-fernandez", "2020")),
    ]
    for opis, redak, ocekivano in JEDINICE:
        kljucevi, _ = AY.extract_biblio_keys(redak)
        check(f"R34: {opis}", ocekivano in kljucevi, (redak[:60], kljucevi))

    # 94 — akronim iz nakladničkog dijela je alias
    for redak, alias, glavni in [
        ("Hrvatski zavod za javno zdravstvo. Atlas. Zagreb: HZJZ, 2024.",
         "hzjz", "hrvatski"),
        ("Dr\u017eavni zavod za statistiku. Istra\u017eivanje. Zagreb: DZS, 2024.",
         "dzs", "dr\u017eavni"),
    ]:
        mapa = AY.biblio_aliasi(redak)
        check(f"R34: akronim {alias.upper()} je alias za {glavni}",
              mapa.get((alias, "2024")) == (glavni, "2024"), mapa)

    # 95 — hrvatski izvor prikaza „(autorski sažetak prema: X, 2026)"
    for tekst, ocekivan in [
        ("(autorski sa\u017eetak prema: Podobnik, 2026)", "podobnik"),
        ("(autorska analiza prema: Porter, 2008)", "porter"),
        ("(Izvor: izrada autora prema: Murray, 2010)", "murray"),
    ]:
        k = C.parse_ay_citation_group(tekst.strip("()"))
        check(f"R34: izvor prikaza daje klju\u010d {ocekivan}",
              any(x[0] == ocekivan for x in k), (tekst, k))
    k = C.parse_ay_citation_group("autorska analiza prema: Porter, 2008")
    check("R34: 'autorska' NIJE klju\u010d",
          not any(x[0].startswith("autorsk") for x in k), k)

    # 96 — razlika izrečena u istoj rečenici je deklarirana
    izlaz = _inv("\u010celik S355 nije kru\u0107i od \u010delika S235. "
                 "Granica popu\u0161tanja iznosi 235 N i 355 N.")
    check("R34: obje vrijednosti u istoj re\u010denici nisu sukob",
          "izre\u010deno je u istoj re\u010denici" in izlaz or "'\u010delik'" not in izlaz,
          izlaz[-350:])

    # 102 — dijalekt navodnika (audit Znahor, B4)
    hr_ihjj = "Tekst \u201eprvi\u201d i \u201edrugi\u201d navod."
    hr_njem = "Tekst \u201eprvi\u201c i \u201edrugi\u201c navod."
    for opis, tekst in [("IHJJ \u201e\u2026\u201d", hr_ihjj), ("njema\u010dki \u201e\u2026\u201c", hr_njem)]:
        n = _tipografija_nalazi(TY, tekst)
        check(f"R35: dosljedan dijalekt ({opis}) NIJE nalaz",
              not any("zatvaraju\u0107eg navodnika" in x for x in n), (opis, n))
    n = _tipografija_nalazi(TY, "Tekst \u201eprvi\u201d i \u201edrugi\u201c navod.")
    check("R35: mije\u0161ana dva oblika JEST nalaz",
          any("dva oblika" in x for x in n), n)

    # 105 — metapodaci (faza A4)
    import os as _os
    import tempfile as _tf
    import zipfile as _zf
    import provjeri_metapodatke as MP
    import check_placeholders as PH2
    from docx import Document as _Doc

    def _s_meta(put, **polja):
        d = _Doc()
        d.add_paragraph("Sveu\u010dili\u0161te u Zagrebu")
        d.add_paragraph("Ana Ani\u0107")
        d.add_heading("1. UVOD", level=1)
        d.add_paragraph("Tekst rada.")
        d.core_properties.author = polja.get("autor", "")
        d.core_properties.last_modified_by = polja.get("zadnji", "")
        d.core_properties.title = polja.get("naslov", "")
        d.save(put)
        return put

    with _tf.TemporaryDirectory() as dd:
        r1 = _s_meta(_os.path.join(dd, "a.docx"),
                     autor="Ana Ani\u0107", zadnji="Provenance Test Generator",
                     naslov="Naslov")
        n1 = MP.nalazi(MP.procitaj(r1), "Ana Ani\u0107")
        poruke = " | ".join(x for _t, x in n1)
        check("R36: trag alata u lastModifiedBy je gre\u0161ka",
              any(t == "g" and "odaje alat" in x for t, x in n1), poruke)

        r2 = _s_meta(_os.path.join(dd, "b.docx"),
                     autor="Student", zadnji="PC", naslov="Naslov")
        n2 = MP.nalazi(MP.procitaj(r2), "Ana Ani\u0107")
        check("R36: ostavljeni predlo\u017eak (Student/PC) je gre\u0161ka",
              sum(1 for t, x in n2 if t == "g" and "predlo\u017eak" in x) == 2,
              [x for t, x in n2])

        r3 = _s_meta(_os.path.join(dd, "c.docx"),
                     autor="Marko Marki\u0107", zadnji="Marko Marki\u0107", naslov="Naslov")
        n3 = MP.nalazi(MP.procitaj(r3), "Ana Ani\u0107")
        check("R36: autor u metapodacima \u2260 naslovnica je gre\u0161ka",
              any("ne govore isto" in x for _t, x in n3), [x for t, x in n3])

        n4 = MP.nalazi(MP.procitaj(r1), "Ana Ani\u0107")
        check("R36: uskla\u0111en autor NIJE nalaz",
              not any("ne govore isto" in x for _t, x in n4), [x for t, x in n4])

        # --postavi piše u NOVU datoteku i ne dira tragove
        izl = _os.path.join(dd, "a-meta.docx")
        MP.postavi(r1, izl, {"zadnji_uredio": "Ana Ani\u0107"}, True)
        check("R36: --postavi ostavlja izvornik netaknut", _os.path.exists(r1))
        check("R36: --postavi pi\u0161e novu datoteku", _os.path.exists(izl))
        m_novi = MP.procitaj(izl)
        check("R36: upisano lastModifiedBy",
              m_novi["core"]["zadnji_uredio"] == "Ana Ani\u0107", m_novi["core"])
        with _zf.ZipFile(r1) as z1, _zf.ZipFile(izl) as z2:
            check("R36: document.xml NIJE diran (rsid i izmjene ostaju)",
                  z1.read("word/document.xml") == z2.read("word/document.xml"))

    # 105b — placeholder predloška: oba smjera
    for tekst, treba in [("[Ime i prezime]", True), ("[Datum obrane]", True),
                         ("[Internet]", False), ("[12, 13]", False), ("[3]", False),
                         ("[zavr\u0161ni rad]", False)]:
        import re as _re
        uz = [u for u, _o in PH2.UZORCI if "ime|prezime" in u][0]
        check(f"R36: placeholder {tekst} \u2192 {treba}",
              bool(_re.search(uz, tekst)) == treba, tekst)

    # 106 — uputnice na prikaze, oba smjera + palatalizacija
    import check_uputnice as CUP
    import tempfile as _tf2
    import os as _os2
    from docx import Document as _D2
    with _tf2.TemporaryDirectory() as dd2:
        def _rad(put, redci):
            d = _D2()
            d.add_heading("1. UVOD", level=1)
            for x in redci:
                d.add_paragraph(x)
            d.save(put)
            return put

        r = _rad(_os2.path.join(dd2, "u1.docx"), [
            "Raspodjela je prikazana na Slici 1.",
            "Slika 1. Raspodjela odgovora",
            "Izvor: autor."])
        a = CUP.analiziraj(r)
        check("R37: palatalizacija (na Slici 1) je uputnica na Sliku 1",
              1 in a["uputnice"].get("Slika", []), a["uputnice"])
        check("R37: spomenut prikaz nije nalaz", not a["nespomenuti"], a["nespomenuti"])

        r = _rad(_os2.path.join(dd2, "u2.docx"), [
            "Tekst bez ijedne uputnice.",
            "Tablica 1. Naslov", "Izvor: autor."])
        a = CUP.analiziraj(r)
        check("R37: nespomenut prikaz JEST nalaz",
              a["nespomenuti"] and a["nespomenuti"][0][0] == "Tablica", a["nespomenuti"])

        r = _rad(_os2.path.join(dd2, "u3.docx"), [
            "Podaci su u Tablici 5.",
            "Tablica 1. Naslov", "Izvor: autor."])
        a = CUP.analiziraj(r)
        check("R37: uputnica na nepostoje\u0107u tablicu JEST nalaz",
              bool(a["u_prazno"]), a["u_prazno"])

        r = _rad(_os2.path.join(dd2, "u4.docx"), [
            "Vidi Tablicu 1 i 2 za usporedbu.",
            "Tablica 1. Prva", "Izvor: autor.",
            "Tablica 2. Druga", "Izvor: autor."])
        a = CUP.analiziraj(r)
        check("R37: raspon 'Tablicu 1 i 2' hvata oba broja",
              a["uputnice"].get("Tablica") == [1, 2], a["uputnice"])

    # 107 — kod 3 je granica, ne pad
    import brojke_iz_rasprave as BR
    with _tf2.TemporaryDirectory() as dd3:
        d = _D2(); d.add_heading("1. UVOD", level=1)
        d.add_paragraph("Rad bez zasebne Rasprave.")
        put = _os2.path.join(dd3, "bez.docx"); d.save(put)
        import io as _io2, contextlib as _cx2
        buf = _io2.StringIO()
        with _cx2.redirect_stdout(buf):
            kod = BR.main(put)
        check("R37: rad bez Rasprave daje kod 3 (granica), ne 2 (pad)",
              kod == 3, (kod, buf.getvalue()[-120:]))

    # 108–110 — statistika, tablice, hipoteze
    import check_statistika as ST
    import check_tablice as TB
    import check_hipoteze as HI
    import tempfile as _t3, os as _o3
    from docx import Document as _D3

    with _t3.TemporaryDirectory() as d3:
        d = _D3(); d.add_heading("5. REZULTATI", level=1)
        d.add_paragraph("Značajnost je utvrđivana na razini p < 0,05. Korišten je t-test.")
        d.add_paragraph("Razlika je statistički značajna (p = 0,322).")
        d.add_paragraph("Razlika nije bila statistički značajna (p = 0,004).")
        d.add_paragraph("Vrijednost je iznosila p = 0,000.")
        d.add_paragraph("Prema spolu razlika nije se statistički značajno "
                        "razlikovala (p = 0,780).")
        d.add_paragraph("Žene su značajno češće birale opciju (p = 0,001).")
        put = _o3.path.join(d3, "s.docx"); d.save(put)
        n = ST.analiziraj(put)
        pv = {x["p"] for x in n["proturjecje"]}
        check("R38: p = 0,322 opisan kao značajan JEST nalaz", "p = 0,322" in pv, pv)
        check("R38: p = 0,004 opisan kao neznačajan JEST nalaz", "p = 0,004" in pv, pv)
        check("R38: p = 0,780 uz negaciju NIJE nalaz", "p = 0,780" not in pv, pv)
        check("R38: p = 0,001 uz tvrdnju NIJE nalaz", "p = 0,001" not in pv, pv)
        check("R38: p = 0,000 je nemoguća vrijednost",
              any("0,000" in x["p"] for x in n["nemoguc_p"]), n["nemoguc_p"])
        check("R38: deklarirani prag se čita iz teksta",
              n["deklarirana_alfa"] == 0.05, n["deklarirana_alfa"])

        d = _D3(); d.add_heading("5. REZULTATI", level=1)
        d.add_paragraph("Tablica 1. Raspodjela ispitanika (n = 100)")
        t = d.add_table(rows=5, cols=3)
        for i, red in enumerate([["Skupina", "n", "Udio (%)"], ["A", "30", "30,0"],
                                 ["B", "25", "25,0"], ["C", "40", "35,0"],
                                 ["Ukupno", "95", "90,0"]]):
            for j, v in enumerate(red):
                t.rows[i].cells[j].text = v
        put = _o3.path.join(d3, "t.docx"); d.save(put)
        r = TB.analiziraj(put)
        vrste = {x["vrsta"] for x in r["nalazi"]}
        check("R39: n iz natpisa naspram zbroja stupca", "n_ne_odgovara" in vrste, r["nalazi"])
        check("R39: postoci koji ne daju 100", "postoci_ne_daju_sto" in vrste, r["nalazi"])

        # „Ukupno" kao NAZIV retka, ne redak zbroja
        d = _D3(); d.add_heading("5. REZULTATI", level=1)
        d.add_paragraph("Tablica 1. Pokazatelji")
        t = d.add_table(rows=4, cols=2)
        for i, red in enumerate([["Pokazatelj", "Iznos"], ["Neto primitak", "4,4"],
                                 ["Ukupno suspendirana sredstva", "32"],
                                 ["Trajno izgubljeno", "2"]]):
            for j, v in enumerate(red):
                t.rows[i].cells[j].text = v
        put = _o3.path.join(d3, "t2.docx"); d.save(put)
        check("R39: 'Ukupno' u nazivu retka NIJE redak zbroja",
              not TB.analiziraj(put)["nalazi"], TB.analiziraj(put)["nalazi"])

        d = _D3(); d.add_heading("3. HIPOTEZE", level=1)
        for x in ("H1. Prva hipoteza.", "H2. Druga hipoteza.", "H3. Treća hipoteza."):
            d.add_paragraph(x)
        d.add_heading("6. RASPRAVA", level=1)
        d.add_paragraph("Hipoteza H1 je prihvaćena na temelju rezultata.")
        d.add_paragraph("H2 se odbacuje jer razlika nije utvrđena.")
        d.add_paragraph("Podaci o H3 prikazani su u Tablici 4.")
        put = _o3.path.join(d3, "h.docx"); d.save(put)
        r = HI.analiziraj(put)
        check("R40: sve tri hipoteze prepoznate (odlomak, ne rečenica)",
              r["postavljene"] == [1, 2, 3], r["postavljene"])
        check("R40: H3 bez presude JEST nalaz", r["bez_presude"] == [3], r["bez_presude"])
        check("R40: H1 i H2 su presuđene", r["presudene"] == [1, 2], r["presudene"])

    # 78 — DOI nije množenje
    n = _tipografija_nalazi(TY, "DOI: 10.1177/1023263X251338198. Dostupno na mre\u017ei.")
    check("R26: DOI s X me\u0111u znamenkama NIJE mno\u017eenje",
          not any("mno\u017eenje" in x for x in n), n)
    n = _tipografija_nalazi(TY, "Uzorak je bio 80 x 80 mm.")
    check("R26: pravo 'x' kao mno\u017eenje JEST nalaz",
          any("mno\u017eenje" in x for x in n), n)



def test_izvori():
    """R27–R29 — veza tvrdnja↔izvor, postojanje jedinice, granica popisa."""
    import os
    import tempfile
    import common as C
    import check_tvrdnja_izvor as TI
    import check_reference_exists as RE
    from docx import Document

    # 79 — popis literature završava na sljedećem naslovu, ne na kraju dokumenta
    body = ("Tijelo rada.\nLiteratura\nHorvat, S. (2018). Knjiga. Zagreb: A.\n"
            "Popis tablica\nTablica 1. Nesto\t10\nSažetak\nRad analizira nesto.")
    lit = C.dio_literature(body)
    check("R27: popis literature staje na 'Popis tablica'",
          "Horvat" in lit and "Tablica 1" not in lit and "analizira" not in lit,
          repr(lit))
    check("R27: bez naslova popisa vraća prazno",
          C.dio_literature("Samo tijelo bez popisa.") == "")

    with tempfile.TemporaryDirectory() as radni:
        izvori = os.path.join(radni, "izvori")
        os.makedirs(izvori)
        with open(os.path.join(izvori, "Closa_2021_x.txt"), "w", encoding="utf-8") as f:
            f.write("Council held 12 hearings. Support stood at 62,1 %.")
        with open(os.path.join(izvori, "Thinus_2025_y.txt"), "w", encoding="utf-8") as f:
            f.write("Compliance improved by 16,3 percentage points.")
        mapa = os.path.join(izvori, "mapa.json")
        with open(mapa, "w", encoding="utf-8") as f:
            f.write('{"closa-2021":{"datoteka":"Closa_2021_x.txt"},'
                    '"thinus-2025":{"datoteka":"Thinus_2025_y.txt"}}')

        d = Document()
        d.add_heading("1. UVOD", level=1)
        d.add_paragraph("Vijeće je održalo 12 saslušanja (Closa, 2021).")
        d.add_paragraph("Usklađenost se popravila za 16,3 postotnih bodova (Closa, 2021).")
        d.add_paragraph("Potpora je iznosila 88,4 % (Thinus, 2025).")
        d.add_heading("Literatura", level=1)
        d.add_paragraph("Closa, Carlos (2021) Institutional logics. JCMS.")
        d.add_paragraph("Thinus, Pauline (2025) Transactional Approach. JCMS.")
        rad = os.path.join(radni, "rad.docx")
        d.save(rad)

        n = TI.provjeri(rad, izvori, mapa)
        check("R28: točno pripisana brojka je potvrđena", n["potvrdeno"] >= 1, n["potvrdeno"])
        krivi = [x["broj"] for x in n["krivi_izvor"]]
        check("R28: brojka iz drugog izvora je PRIPISANO KRIVOM IZVORU",
              "16,3" in krivi, n["krivi_izvor"])
        nema = [x["broj"] for x in n["nije_nadeno"]]
        check("R28: brojke koje nema nigdje su NIJE NAĐENO", "88,4" in nema, n["nije_nadeno"])

        r = RE.provjeri(rad, izvori, mapa)
        check("R29: jedinica s priloženom građom nije nepotvrđena",
              len(r["nepotvrdeno"]) == 0, r)

    # 80 — službena oznaka akta u više oblika
    for akt in ["Uredba (EU, Euratom) 2020/2092 Europskog parlamenta i Vijeća od 2020.",
                "Zakon o radu, NN 93/14, 2014.",
                "Presuda Suda EU, ECLI:EU:C:2024:493, 2024."]:
        check(f"R29: službena oznaka prepoznata: {akt[:28]}",
              bool(RE.SLUZBENI.search(akt)), akt)

    # ISBN kontrolna znamenka
    check("R29: valjan ISBN-13 prolazi", RE._isbn_valjan("978-3-16-148410-0"))
    check("R29: neispravan ISBN-13 pada", not RE._isbn_valjan("978-3-16-148410-1"))


def _tipografija_nalazi(TY, tekst):
    """Pokreni tipografske provjere nad golim tekstom, bez .docx-a."""
    import io
    import contextlib
    from unittest import mock
    buf = io.StringIO()
    with mock.patch.object(TY, "load_docx_text", lambda p, include_tables=True: (tekst, [], None)):
        with contextlib.redirect_stdout(buf):
            TY.main("(test)")
    return [r.strip() for r in buf.getvalue().splitlines() if "⚠" in r]


def test_manifest():
    """Kvar 1/11: manifest motora mora se slagati s otiskom koda."""
    import subprocess
    r = subprocess.run([sys.executable, os.path.join(SCRIPTS, "osvjezi_contract.py")],
                       capture_output=True, text=True)
    check("R22: engine_contract.json se slaže s otiskom koda "
          "(popravak: python3 osvjezi_contract.py --upisi)",
          r.returncode == 0, r.stdout.strip().splitlines()[-1:] )

    # R23: git s core.autocrlf=true (zadano na Windowsu) piše CRLF u radno stablo.
    # Dok je otisak hashirao sirove bajtove, isti je commit davao dva otiska, pa je
    # manifest bio ispravan na jednoj platformi i pogrešan na drugoj — R22 bi tada
    # padao ovisno o tome tko ga pokreće, a ne o tome slažu li se kod i ugovor.
    import importlib.util
    import shutil
    otisci = {}
    for ime, pretvori in (("crlf", lambda d: d.replace(b"\r\n", b"\n").replace(b"\n", b"\r\n")),
                          ("lf", lambda d: d.replace(b"\r\n", b"\n"))):
        t = tempfile.mkdtemp(prefix="otisak_%s_" % ime)
        for n in os.listdir(SCRIPTS):
            if n.endswith(".py"):
                with open(os.path.join(SCRIPTS, n), "rb") as fh:
                    open(os.path.join(t, n), "wb").write(pretvori(fh.read()))
        sp = importlib.util.spec_from_file_location("ka_" + ime,
                                                    os.path.join(t, "katedra_adapter.py"))
        m = importlib.util.module_from_spec(sp)
        sys.path.insert(0, t)
        try:
            sp.loader.exec_module(m)
            otisci[ime] = m.otisak_motora()
        finally:
            sys.path.pop(0)
            shutil.rmtree(t, ignore_errors=True)
    check("R41: otisak motora je isti za CRLF i LF radno stablo",
          otisci.get("crlf") == otisci.get("lf"), otisci)


def main():
    tmp = tempfile.mkdtemp(prefix="rad_audit_fixtures_")
    fx = os.path.join(tmp, "fixtures")
    make_fixtures.build(fx)
    src = os.path.join(fx, "izvori")

    # --- 1) apply_safe_fixes: navodnici po odlomku, inč-oznaka preskočena, warning za nesparen ---
    import apply_safe_fixes
    out = os.path.join(tmp, "quotes_fixed.docx")
    txt, code = capture(apply_safe_fixes.main,
                         [os.path.join(fx, "quotes.docx"), out, "--no-mult", "--no-breaks", "--no-autofit"])
    from docx import Document
    d = Document(out)
    texts = [p.text for p in d.paragraphs]
    check("quotes: par ispravno pretvoren", texts[0] == 'On je rekao „pozdrav” i otišao.', texts[0])
    check("quotes: inč-oznaka NETAKNUTA", '12"' in texts[1], texts[1])
    check("quotes: nesparen odlomak i dalje ima jedan „", texts[2].count("„") == 1, texts[2])
    check("quotes: razdvojen kroz 3 runa ispravno spojen/pretvoren",
          texts[3] == 'Prvi „razdvojen navodnik” u tri runa.', texts[3])
    check("quotes: upozorenje za neparan odlomak ispisano", "neparan broj navodnika" in txt)

    # --- 2) check_citations: numerirani naslov LITERATURE prepoznat, bez lažnih nalaza ---
    import check_citations
    txt, code = capture(check_citations.main, os.path.join(fx, "ieee_numbered_heading.docx"))
    check("check_citations: numerirani naslov prepoznat (bez rupa/siročadi)",
          "rupe u numeraciji: nema" in txt and "SIROČAD" in txt and code == 0, txt)

    # --- 3) check_citations_authoryear: nalazi siroče + citat bez reference ---
    import check_citations_authoryear
    txt, code = capture(check_citations_authoryear.main, os.path.join(fx, "author_year.docx"))
    check("authoryear: nalazi siroče horvat/2018", "'horvat', '2018'" in txt, txt)
    check("authoryear: nalazi citat bez reference kovač/2021", "'kovač', '2021'" in txt, txt)
    check("authoryear: exit code != 0 (ima nalaza)", code != 0)

    # --- 4) common.detect_citation_style ---
    from common import detect_citation_style, load_docx_text
    body, cells, _ = load_docx_text(os.path.join(fx, "author_year.docx"), include_tables=True)
    style, counts = detect_citation_style(body + "\n" + "\n".join(cells))
    check("detect_citation_style: prepoznaje authoryear", style == "authoryear", (style, counts))

    # --- 5) common.load_supplementary_text: fusnota se čita ---
    from common import load_supplementary_text
    sup = load_supplementary_text(os.path.join(fx, "footnote.docx"))
    check("load_supplementary_text: fusnota pročitana", "Ivić, 2020" in sup["footnotes"], sup)

    # --- 6) check_fields: neprihvaćene izmjene detektirane ---
    import check_fields
    txt, code = capture(check_fields.main, os.path.join(fx, "tracked_changes.docx"))
    check("check_fields: w:ins detektiran", "NEPRIHVAĆENE IZMJENE" in txt and code != 0, txt[:200])
    txt2, code2 = capture(check_fields.main, os.path.join(fx, "quotes.docx"))
    check("check_fields: bez tracked changes -> nema upozorenja", "NEPRIHVAĆENE IZMJENE" not in txt2 and code2 == 0)

    # --- 7) domains: auto-detekcija elektro ---
    from domains import detect_domain
    body, cells, _ = load_docx_text(os.path.join(fx, "elektro.docx"), include_tables=True)
    dom, scores = detect_domain(body + "\n" + "\n".join(cells))
    check("detect_domain: elektro prepoznat", dom == "elektro", (dom, scores))

    # --- 8) numbers_inventory: koristi elektro domenu ---
    import numbers_inventory
    txt, code = capture(numbers_inventory.main, os.path.join(fx, "elektro.docx"))
    check("numbers_inventory: domena elektro u ispisu", "elektro" in txt, txt[:200])

    # --- 9) cross_check: kontekst prikazan, lažni pozitivac vidljiv u kontekstu ---
    import cross_check
    txt, code = capture(cross_check.main, [os.path.join(fx, "cross_fp.docx"), src])
    check("cross_check: prikazuje kontekst (└─)", "└─" in txt, txt[:300])
    check("cross_check: lažni pozitivac '40 t' vidljiv u kontekstu s 'tvrtke'",
          any("tvrtke" in line for line in txt.split("\n") if "40" in line or "└─" in line), txt)

    # --- 10) check_overlap: neoznačeno vs označeno vs parafraza ---
    import check_overlap
    txt, code = capture(check_overlap.main, [os.path.join(fx, "overlap.docx"), src])
    check("check_overlap: odlomak #1 (neoznačeno) flagiran",
          "#1" in txt and "NEMA vidljive oznake" in txt, txt[:400])
    check("check_overlap: odlomak #2 (označeno) prepoznat kao izgleda označeno",
          "izgleda označeno" in txt, txt[:600])
    check("check_overlap: odlomak #3 (parafraza) NIJE flagiran",
          "#3" not in txt.split("SAŽETAK")[0])

    # --- 11) common.sentences: kratice ne razbijaju rečenicu ---
    from common import sentences
    s = sentences("Firma d.o.o. posluje dobro. Vidi npr. sljedeći primjer.")
    check("sentences: 'd.o.o.' ne razbija rečenicu", s[0] == "Firma d.o.o. posluje dobro.", s)
    check("sentences: 'npr.' ne razbija rečenicu", s[1] == "Vidi npr. sljedeći primjer.", s)

    # =====================================================================
    # Runda 2 — regresijski testovi za nalaze drugog audita (17 fixeva)
    # =====================================================================
    import zipfile as _zip

    # --- R1) citat samo u ćeliji tablice NIJE siroče; [2020] nije citat ---
    txt, code = capture(check_citations.main, os.path.join(fx, "table_cite.docx"))
    check("R1 tablica: [2] iz ćelije NIJE siroče", "SIROČAD (u popisu, ne citirano): nema" in txt, txt)
    check("R1 godina: [2020] prijavljen kao godina, ne citat",
          "vjerojatno godina" in txt and "[2020]" in txt, txt)
    check("R1: interno konzistentno (exit 0)", code == 0)

    # --- R2) autor-godina citat SAMO u fusnoti se vidi (end-to-end) ---
    txt, code = capture(check_citations_authoryear.main, os.path.join(fx, "fn_ay_cite.docx"))
    check("R2 fusnota: citat iz fusnote prebrojan", "Citirano u tekstu (uklj. fusnote/endnote): 1" in txt, txt)
    check("R2 fusnota: nema lažnog siročeta", "SIROČAD (u popisu, ne citirano): nema" in txt, txt)

    # --- R2b) izvori bez osobnog autora + signalna riječ/napomena u citatu ---
    txt, code = capture(check_citations_authoryear.main,
                        os.path.join(fx, "author_year_institutions.docx"))
    check("R2b institucije/mediji: svih 5 jedinica prepoznato",
          "Definirano u popisu literature (prezime+godina): 5" in txt, txt)
    check("R2b institucije/mediji: nema lažnog siročeta",
          "SIROČAD (u popisu, ne citirano): nema" in txt, txt)
    check("R2b institucije/mediji: nema lažnog citata bez reference",
          "CITAT BEZ REFERENCE: nema" in txt, txt)
    check("R2b institucije/mediji: konzistentan dokument prolazi", code == 0, txt)

    # Guardovi: pravi manjak i pogrešna godina i dalje moraju ostati nalazi.
    sir, bez = check_citations_authoryear.uskladi_kljuceve(
        {("danas.hr", "2025"), ("unesco", "3035")},
        {("unesco", "2021")})
    check("R2b guard: stvarno nedostajući medij ostaje citat bez reference",
          ("danas.hr", "2025") in bez, (sir, bez))
    check("R2b guard: nemoguća/pogrešna godina ne spaja se s pravom",
          ("unesco", "3035") in bez and ("unesco", "2021") in sir, (sir, bez))

    # --- R3) --no-indent ne kontaminira rPr; XML valjan raspored ---
    out_r3 = os.path.join(tmp, "rpr_fixed.docx")
    txt, code = capture(apply_safe_fixes.main,
                         [os.path.join(fx, "rpr_spacing.docx"), out_r3, "--no-indent", "--no-quotes", "--no-mult"])
    import re as _re
    x = _zip.ZipFile(out_r3).read("word/document.xml").decode()
    check("R3: NEMA w:after u rPr spacingu", not _re.search(r"<w:rPr>[^<]*<w:spacing[^>]*w:after", x), None)
    check("R3: w:after JEST u pPr", bool(_re.search(r'<w:pPr>.*?<w:spacing[^>]*w:after="120"', x, _re.S)), None)

    # --- R4) engleski “…” par ostaje; njemački „…“ se popravi ---
    out_r4 = os.path.join(tmp, "eng_fixed.docx")
    capture(apply_safe_fixes.main,
            [os.path.join(fx, "eng_quotes.docx"), out_r4, "--no-mult", "--no-breaks", "--no-autofit"])
    from docx import Document as _D
    pts = [p.text for p in _D(out_r4).paragraphs]
    check("R4: engleski par netaknut", "“smart control”" in pts[0], pts[0])
    check("R4: njemački par popravljen u „…”", "„citat”" in pts[1], pts[1])

    # --- R5) hex literal preskočen, prava multiplikacija pretvorena ---
    out_r5 = os.path.join(tmp, "hex_fixed.docx")
    capture(apply_safe_fixes.main,
            [os.path.join(fx, "hex.docx"), out_r5, "--no-quotes", "--no-breaks", "--no-autofit"])
    t5 = _D(out_r5).paragraphs[0].text
    check("R5: 0x41/0xFF00 netaknuti", "0x41" in t5 and "0xFF00" in t5, t5)
    check("R5: 80x80 → 80 × 80", "80 × 80" in t5, t5)

    # --- R6) _normal_firstline0 NE dira tuđi stil ---
    out_r6 = os.path.join(tmp, "styles_fixed.docx")
    capture(apply_safe_fixes.main,
            [os.path.join(fx, "styles_extra.docx"), out_r6, "--no-indent", "--no-quotes", "--no-mult"])
    s6 = _zip.ZipFile(out_r6).read("word/styles.xml").decode()
    m6 = _re.search(r'styleId="Citat9".*?</w:style>', s6, _re.S)
    check("R6: Citat9 firstLine=709 netaknut", m6 and 'firstLine="709"' in m6.group(0), None)

    # --- R7) tekst POSLIJE inline textboxa se obrađuje (ne preskače tiho) ---
    out_r7 = os.path.join(tmp, "txbx_fixed.docx")
    capture(apply_safe_fixes.main,
            [os.path.join(fx, "txbx.docx"), out_r7, "--no-mult", "--no-breaks", "--no-autofit"])
    x7 = _zip.ZipFile(out_r7).read("word/document.xml").decode()
    t7 = _re.findall(r"<w:t[^>]*>([^<]*)</w:t>", x7)
    check("R7: navodnici PRIJE textboxa pretvoreni", any("„prvi citat”" in t for t in t7), t7[:3])
    check("R7: navodnici POSLIJE textboxa pretvoreni", any("„drugi citat”" in t for t in t7), t7[:3])

    # --- R8) %/°/V/Hz jedinice + detekcija sukoba vrijednosti ---
    txt, code = capture(numbers_inventory.main, os.path.join(fx, "elektro_konflikt.docx"))
    check("R8: V i Hz u inventaru", "V    :" in txt.replace("V   :", "V    :") or " V " in txt, txt)
    check("R8: % u inventaru", _re.search(r"%\s*:", txt) is not None, txt)
    check("R8: ° u inventaru", _re.search(r"°\s*:", txt) is not None, txt)
    check("R8: sukob 'napon' 230 vs 400 flagiran",
          "'napon'" in txt and "230" in txt and "400" in txt and "⚠" in txt, txt)
    check("R8: sekcija rečenica NIJE prazna (skriveni filter maknut)",
          "Sustav radi na naponu" in txt, txt)

    # --- R9) mixed stil: generate_report pokreće OBA B checkera + napomenu ---
    import generate_report as _gr
    out_j9 = os.path.join(tmp, "mixed.json")
    capture(_gr.main, [os.path.join(fx, "mixed_style.docx"), "--out",
                        os.path.join(tmp, "mixed.md"), "--json", out_j9])
    import json as _json
    p9 = _json.load(open(out_j9, encoding="utf-8"))
    b_phases = [k for k in p9["phase_exit_codes"] if k.startswith("B")]
    check("R9: mixed → IEEE i autor-godina faze + napomena",
          len(b_phases) == 3 and any("Napomena" in k for k in b_phases), b_phases)

    # --- R10) cross_check exit kod 1 kad ima nepotvrđenih tvrdnji ---
    # (elektro_konflikt ima 230 V koji NE postoji u izvorima → mora dati exit 1;
    #  cross_fp fixture namjerno ima sve "pronađeno" pa nije prikladan ovdje)
    txt, code = capture(cross_check.main, [os.path.join(fx, "elektro_konflikt.docx"), src])
    check("R10: cross_check exit != 0 uz promašaje",
          code != 0 and "nije nađeno u izvorima" in txt, (code, txt[-200:]))

    # --- R11) '-ost.' kraj rečenice se NE guta ('st.' kratica) ---
    s11 = sentences("Provjerena je nosivost. Rezultat je dobar.")
    check("R11: 'nosivost.' ispravno završava rečenicu", len(s11) == 2, s11)

    # --- R12) golo 'S' nije claim; oznaka čelika s kvalitetom jest ---
    from cross_check import auto_claims as _ac
    cl12, _dom = _ac("Čelik S 355 i stup tipa S te kvaliteta S355J2 u konstrukciji od čelika s vijkom i pločom uz profil.")
    check("R12: golo 'S' NIJE claim", "S" not in cl12, cl12)
    check("R12: 'S 355' i 'S355J2' jesu claimovi", "S 355" in cl12 and "S355J2" in cl12, cl12)

    # --- 12) generate_report: end-to-end, kritični nalazi bucketirani ---
    import generate_report
    out_md = os.path.join(tmp, "report.md")
    out_json = os.path.join(tmp, "report.json")
    txt, code = capture(generate_report.main,
                         [os.path.join(fx, "author_year.docx"), "--sources", src,
                          "--out", out_md, "--json", out_json])
    check("generate_report: .md spremljen", os.path.exists(out_md))
    check("generate_report: .json spremljen", os.path.exists(out_json))
    import json
    with open(out_json, encoding="utf-8") as f:
        payload = json.load(f)
    check("generate_report: siroče/citat-bez-reference u KRITIČNO bucketu",
          payload["counts"]["kritično"] >= 2, payload["counts"])

    shutil.rmtree(tmp, ignore_errors=True)

    # --- R13: zakrpe nađene na obranjenom FPZG radu (kolovoz 2026.) ---
    # Svih pet ima isti oblik: alat je bio kalibriran na jedan citatni dijalekt, pa
    # rad koji radi nešto drukčije — ali ispravno — prijavljuje kao pogrešan. To je
    # najskuplja vrsta kvara ovdje: student dobije popis nepostojećih grešaka i,
    # ako mu vjeruje, pokvari rad koji je bio dobar.
    import common as C13
    import check_citations_authoryear as B13

    def _zagradni(t):
        nadjeno = C13.CITE_AY_RE.findall(t)
        return C13.parse_ay_citation_group(nadjeno[0]) if nadjeno else set()

    # (1) lokator stranice iza godine — FPZG Upute propisuju baš taj oblik
    check("R13: (Becker, 2007: 45) je citat",
          ("becker", "2007") in _zagradni("(Becker, 2007: 45)"))
    check("R13: (Streeck, 2014: xiv) je citat",
          ("streeck", "2014") in _zagradni("(Streeck, 2014: xiv)"))
    check("R13: narativni lokator (Krippner (2005: 174))",
          ("krippner", "2005") in C13.parse_ay_narrative("Krippner (2005: 174) tvrdi"))

    # (2) sufiks godine je dio identiteta, inače se dva rada slijevaju u jedan ključ
    check("R13: sufiks 2013a/2013b preživi zagradni oblik",
          C13.parse_ay_segment("Becker, 2013a") == ("becker", "2013a")
          and C13.parse_ay_segment("Becker, 2013b") == ("becker", "2013b"))

    # (3) popis literature u hrvatskim/FPZG oblicima
    for redak, kljuc in [
        ("Becker, Gary (2007) Ekonomski pristup. Zagreb: Naklada.", ("becker", "2007")),
        ("Becker, G. (2007) Ekonomski pristup. Zagreb: Naklada.", ("becker", "2007")),
        ("Van der Zwan, Natascha (2014) Making sense. SER 12(1).", ("zwan", "2014")),
        ("HNB (Hrvatska narodna banka) (2023) Izvješće. Zagreb: HNB.", ("hnb", "2023")),
        ("easyJet plc (2025.), Full year results, Luton: easyJet plc", ("easyjet", "2025")),
        ("Podravka d.d. (2024.) Izvješće. Koprivnica: Podravka.", ("podravka", "2024")),
    ]:
        kljucevi, _ = B13.extract_biblio_keys(redak)
        check(f"R13: popis literature — {redak[:34]}", kljuc in kljucevi, kljucevi)

    check("R13: obična rečenica u popisu NIJE referenca",
          B13.extract_biblio_keys(
              "ovo je obična rečenica koja spominje 2024. godinu") == (set(), 1))

    # (4) tekst i popis moraju dati ISTI ključ za istu referencu
    iz_teksta = _zagradni("(Van der Zwan, 2014)")
    iz_popisa, _ = B13.extract_biblio_keys("Van der Zwan, Natascha (2014) Making sense.")
    check("R13: Van der Zwan daje isti ključ iz teksta i iz popisa",
          bool(iz_teksta & iz_popisa), (iz_teksta, iz_popisa))

    # (5) institucija s malom riječi u imenu, uz granicu da proza ne postane citat
    check("R13: Europska komisija (2021.) je citat",
          ("europska", "2021") in C13.parse_ay_narrative("Europska komisija (2021.) je odobrila"))
    check("R13: proza sa zagradnom godinom NIJE citat",
          C13.parse_ay_narrative(
              "Analiza je provedena u promatranom razdoblju (2021.)") == set())

    # (6) kratica institucije kao ALIAS, ne kao drugi unos u popisu
    #     Pravilo 34: uz tri slučaja koji moraju proći ide i onaj koji mora pasti,
    #     inače se ne zna razlikuje li alias išta ili samo gasi nalaze.
    for opis, redak, alias, glavni in [
        ("pun naziv pa kratica",
         "Hrvatska narodna banka (HNB) (2023). Bilten o bankama.", "hnb", "hrvatska"),
        ("kratica pa pun naziv",
         "HNB (Hrvatska narodna banka) (2023). Bilten o bankama.", "hrvatska", "hnb"),
    ]:
        mapa = B13.biblio_aliasi(redak)
        check(f"R13: alias institucije — {opis}",
              mapa.get((alias, "2023")) == (glavni, "2023"), mapa)

    kljucevi, _ = B13.extract_biblio_keys(
        "Hrvatska narodna banka (HNB) (2023). Bilten o bankama.")
    check("R13: alias NE stvara drugi unos u popisu (inače lažno siroče)",
          len(kljucevi) == 1 and ("hrvatska", "2023") in kljucevi, kljucevi)

    check("R13: zagrada bez slova nije alias",
          B13.biblio_aliasi("Zavod za statistiku (2022) (2022). Ljetopis.") == {},
          B13.biblio_aliasi("Zavod za statistiku (2022) (2022). Ljetopis."))

    check("R13: osobni autor nema alias",
          B13.biblio_aliasi("Becker, Gary (2007) Ekonomski pristup.") == {})

    test_r16_vancouver()
    test_r14_r15()
    test_v195()
    test_stvarni_rad()
    test_izvori()
    test_manifest()

    # --- report ---
    passed = sum(1 for _, ok, _ in RESULTS if ok)
    print("=" * 70)
    print(f"REZULTATI TESTOVA: {passed}/{len(RESULTS)} prošlo")
    print("=" * 70)
    for name, ok, detail in RESULTS:
        mark = "✓" if ok else "✗ FAIL"
        print(f"  {mark:8} {name}")
        if not ok and detail:
            print(f"           detalj: {detail}")
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
