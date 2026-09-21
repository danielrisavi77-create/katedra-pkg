#!/usr/bin/env python3
"""Independent release-review regressions for the 2026-09-20 Katedra patch.

Usage: python -B rad-audit/scripts/tests/test_zakrpa_release.py
Uses real package imports and synthetic temporary DOCX fixtures.
No private manuscript, Claude call, or installation mutation is required.
"""
from __future__ import annotations
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from docx import Document

ROOT = Path(__file__).resolve().parents[3]
SCRIPTS = ROOT / 'rad-audit' / 'scripts'
sys.path.insert(0, str(SCRIPTS))


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f'Cannot load {path}')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


CC = load_module('review_check_citations', SCRIPTS / 'check_citations.py')
INV = load_module('review_inventory', SCRIPTS / 'inventar_tvrdnji.py')
COMMON = load_module('review_common', SCRIPTS / 'common.py')
HKS_SOURCE = ROOT / 'katedra-lite/scripts/provjeri_hks_fzs.py'
HKS = load_module('review_hks', HKS_SOURCE)


GR = load_module("review_generate_report", SCRIPTS / "generate_report.py")
PHASE_H = GR._faza_upute
REFS = [f'{i}. Autor A. Naslov {i}. Zagreb; 2020.' for i in range(1, 4)]
FRONT = ['TEMELJNA DOKUMENTACIJSKA KARTICA', 'BASIC DOCUMENTATION CARD',
         'SAŽETAK', 'SUMMARY', 'Sadržaj', 'POPIS KRATICA']


def make_docx(path: Path, paragraphs: list[str], refs: list[str] | None = None) -> str:
    d = Document()
    for text in paragraphs:
        d.add_paragraph(text)
    if refs is not None:
        d.add_heading('LITERATURA', level=1)
        for text in refs:
            d.add_paragraph(text)
    d.save(path)
    return str(path)


class ReleaseRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.dir = Path(self.tmp.name)

    def citation_result(self, extra=''):
        path = make_docx(self.dir / 'citations.docx',
                         ['1. UVOD', 'Skrb je opisana (1) i drugdje (2,3).', extra], REFS)
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = CC.main(path, 'vancouver')
        return code, buf.getvalue()

    def test_valid_vancouver_document_passes(self):
        code, output = self.citation_result()
        self.assertEqual(code, 0, output)

    def test_undefined_single_reference_is_blocked(self):
        code, output = self.citation_result('Dodatni dokaz nalazi se u izvoru (99).')
        self.assertNotEqual(code, 0, output)
        self.assertIn('⚠ CITAT BEZ REFERENCE', output)

    def test_undefined_group_reference_is_not_silently_approved(self):
        code, output = self.citation_result('Dokaz potvrđuju izvori (1,99).')
        self.assertNotEqual(code, 0, output)

    def test_undefined_reference_range_is_not_silently_approved(self):
        code, output = self.citation_result('Dokaz potvrđuju izvori (1–5).')
        self.assertNotEqual(code, 0, output)

    def test_statistical_symbol_remains_a_non_citation(self):
        self.assertEqual(COMMON.find_vancouver_citations(
            'Welchov t(106,08) = −1,32; p = 0,190'), [])

    def test_ieee_inventory_reads_standard_bracketed_bibliography(self):
        path = make_docx(self.dir / 'ieee.docx',
                         ['Uzorak sadrži 42 ispitanika [1].'],
                         ['[1] Autor A. Naslov. Zagreb; 2020.'])
        result, error = INV.analiza(path)
        self.assertIsNone(error, error)
        self.assertTrue(result['tvrdnje'][1]['brojcane'])

    def test_vancouver_inventory_still_reads_numbered_bibliography(self):
        path = make_docx(self.dir / 'vancouver.docx',
                         ['Uzorak sadrži 42 ispitanika (1).'], [REFS[0]])
        result, error = INV.analiza(path)
        self.assertIsNone(error, error)
        self.assertTrue(result['tvrdnje'][1]['brojcane'])

    def non_hks_doc(self):
        return make_docx(self.dir / 'fpzg.docx',
                         ['Fakultet političkih znanosti Sveučilišta u Zagrebu.',
                          '1. UVOD', 'Ovo nije HKS-FZS dokument.'])

    def call_phase(self, profile=None):
        previous = os.environ.pop('KATEDRA_LITE', None)
        try:
            return PHASE_H(self.non_hks_doc(), profile)
        finally:
            if previous is not None:
                os.environ['KATEDRA_LITE'] = previous

    def test_unspecified_faculty_does_not_apply_hks_rules(self):
        _name, output, _code = self.call_phase()
        self.assertNotIn('HKS-FZS UPUTE —', output, output)

    def test_missing_explicit_checker_does_not_fall_back_to_hks(self):
        _name, output, _code = self.call_phase(str(self.dir / 'missing-fpzg-checker.py'))
        self.assertNotIn('HKS-FZS UPUTE —', output, output)

    def test_explicit_checker_is_executed(self):
        checker = self.dir / 'explicit_checker.py'
        checker.write_text("print('EXPLICIT_CHECKER_EXECUTED')\n", encoding='utf-8')
        _name, output, code = self.call_phase(str(checker))
        self.assertEqual(code, 0, output)
        self.assertIn('EXPLICIT_CHECKER_EXECUTED', output)

    @staticmethod
    def order_finding(result):
        return next(x for x in result['nalazi']
                    if x['poruka'].startswith('redoslijed dijelova'))

    def test_missing_structure_is_not_a_measured_order_pass(self):
        result = HKS.analiza(self.non_hks_doc())
        finding = self.order_finding(result)
        self.assertIsNot(finding['ok'], True, finding)

    def test_original_correct_frontmatter_order_passes(self):
        path = make_docx(self.dir / 'ordered.docx', FRONT + ['1. UVOD', 'Tekst uvoda.'])
        self.assertIs(self.order_finding(HKS.analiza(path))['ok'], True)

    def test_original_wrong_frontmatter_order_is_rejected(self):
        swapped = FRONT[:4] + [FRONT[5], FRONT[4]]
        path = make_docx(self.dir / 'swapped.docx', swapped + ['1. UVOD', 'Tekst uvoda.'])
        self.assertIs(self.order_finding(HKS.analiza(path))['ok'], False)

    def test_order_is_read_from_profile_not_a_duplicate_constant(self):
        # Mutation experiment only. This does not change real faculty rules.
        isolated = self.dir / 'isolated/katedra-lite'
        target = isolated / 'scripts/provjeri_hks_fzs.py'
        profile = isolated / 'references/fakulteti/hks-fzs.json'
        target.parent.mkdir(parents=True)
        profile.parent.mkdir(parents=True)
        shutil.copyfile(HKS_SOURCE, target)
        data = json.loads((ROOT / 'katedra-lite/references/fakulteti/hks-fzs.json')
                          .read_text(encoding='utf-8'))
        sections = data['struktura']['opseg']['diplomski']['dijelovi']
        sections['sadrzaj']['redoslijed'], sections['kratice']['redoslijed'] = 9, 8
        profile.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
        module = load_module('review_hks_profile_mutation', target)
        swapped = FRONT[:4] + [FRONT[5], FRONT[4]]
        path = make_docx(self.dir / 'profile_order.docx', swapped + ['1. UVOD', 'Tekst uvoda.'])
        finding = self.order_finding(module.analiza(path))
        self.assertIs(finding['ok'], True, finding)


    def test_undefined_reference_cannot_be_disguised_as_decimal(self):
        code, output = self.citation_result('Dokaz je opisan (15,96).')
        self.assertNotEqual(code, 0, output)
        self.assertIn('⚠ CITAT BEZ REFERENCE', output)

    def test_unknown_faculty_is_explicitly_unmeasured(self):
        _name, output, code = self.call_phase()
        self.assertEqual(code, 3, output)
        self.assertIn('NE tvrdi', output)

    def test_invalid_explicit_checker_is_an_error(self):
        _name, output, code = self.call_phase(str(self.dir / 'missing.py'))
        self.assertEqual(code, 2, output)
        self.assertIn('missing.py', output)

    def test_environment_is_not_faculty_consent(self):
        from unittest.mock import patch
        with patch.dict(os.environ, {'KATEDRA_LITE': str(ROOT / 'katedra-lite')}):
            _name, output, code = PHASE_H(self.non_hks_doc())
        self.assertEqual(code, 3, output)
        self.assertNotIn('HKS-FZS UPUTE —', output)

    def test_checker_stderr_is_preserved_with_stdout(self):
        checker = self.dir / 'diagnostic.py'
        checker.write_text("import sys\nprint('OUT_MARKER')\nprint('ERR_MARKER', file=sys.stderr)\nsys.exit(2)\n", encoding='utf-8')
        _name, output, code = self.call_phase(str(checker))
        self.assertEqual(code, 2, output)
        self.assertIn('OUT_MARKER', output)
        self.assertIn('ERR_MARKER', output)

    def test_checker_timeout_is_not_a_declared_profile_boundary(self):
        from unittest.mock import patch
        checker = self.dir / 'slow.py'
        checker.write_text('print("not executed")\n', encoding='utf-8')
        with patch('subprocess.run', side_effect=subprocess.TimeoutExpired(['checker'], 120)):
            _name, output, code = self.call_phase(str(checker))
        self.assertEqual(code, 2, output)

    def test_adapter_preserves_unmeasured_phase_h(self):
        path = self.non_hks_doc()
        output = self.dir / 'adapter.json'
        result = subprocess.run(
            [sys.executable, '-B', str(SCRIPTS / 'katedra_adapter.py'), path,
             '--json', str(output), '--out', str(self.dir / 'report.md')],
            capture_output=True, text=True, encoding='utf-8', timeout=60)
        self.assertTrue(output.is_file(), result.stdout + result.stderr)
        data = json.loads(output.read_text(encoding='utf-8'))
        self.assertEqual(data['contract_version'], '1')
        self.assertEqual(data['phase_exit_codes']['H — Upute fakulteta'], 3)
        self.assertTrue(any(x['phase'] == 'H — Upute fakulteta' for x in data['findings']))
        self.assertNotIn('HKS-FZS UPUTE —', (self.dir / 'report.md').read_text(encoding='utf-8'))

    def test_adapter_preserves_selected_checker_failure(self):
        checker = self.dir / 'selected.py'
        checker.write_text("print('SELECTED_FACULTY_FAILURE')\nraise SystemExit(2)\n", encoding='utf-8')
        output = self.dir / 'selected.json'
        result = subprocess.run(
            [sys.executable, '-B', str(SCRIPTS / 'katedra_adapter.py'), self.non_hks_doc(),
             '--profil', str(checker), '--json', str(output), '--out', str(self.dir / 'selected.md')],
            capture_output=True, text=True, encoding='utf-8', timeout=60)
        self.assertNotEqual(result.returncode, 0)
        self.assertTrue(output.is_file(), result.stdout + result.stderr)
        data = json.loads(output.read_text(encoding='utf-8'))
        self.assertEqual(data['phase_exit_codes']['H — Upute fakulteta'], 2)
        self.assertTrue(any(x['phase'] == 'H — Upute fakulteta' and x['severity'] == 'kritično'
                            for x in data['findings']))
        self.assertIn('SELECTED_FACULTY_FAILURE', (self.dir / 'selected.md').read_text(encoding='utf-8'))

    def test_unknown_order_has_no_green_checkmark(self):
        result = HKS.analiza(self.non_hks_doc())
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            HKS.ispisi(result)
        line = next(x for x in buf.getvalue().splitlines() if 'redoslijed dijelova:' in x)
        self.assertNotIn('✅', line)
        self.assertIn('neizmjeren', buf.getvalue().lower())

    def test_unknown_order_cli_is_not_success(self):
        with contextlib.redirect_stdout(io.StringIO()):
            code = HKS.main([self.non_hks_doc()])
        self.assertEqual(code, 3)

    def test_missing_profile_does_not_reuse_hardcoded_order(self):
        target = self.dir / 'without_profile/scripts/provjeri_hks_fzs.py'
        target.parent.mkdir(parents=True)
        shutil.copyfile(HKS_SOURCE, target)
        module = load_module('hks_missing_profile', target)
        path = make_docx(self.dir / 'without_profile.docx', FRONT + ['1. UVOD'])
        finding = self.order_finding(module.analiza(path))
        self.assertIsNone(finding['ok'], finding)

    def test_hks_table_reference_does_not_cross_paragraph_boundary(self):
        path = make_docx(self.dir / 'table.docx', [
            'Obilježja su prikazana u Tablici 1.', 'n – broj ispitanika; % – postotak'])
        finding = next(x for x in HKS.analiza(path)['nalazi']
                       if x['poruka'].startswith('upućivanje „Tablica N.'))
        self.assertTrue(finding['ok'], finding)

    def test_real_plural_table_punctuation_is_still_a_finding(self):
        path = make_docx(self.dir / 'plural.docx', ['Rezultati u Tablicama 7. i 8. pokazuju razliku.'])
        finding = next(x for x in HKS.analiza(path)['nalazi']
                       if x['poruka'].startswith('upućivanje „Tablica N.'))
        self.assertFalse(finding['ok'], finding)

    def test_statistics_and_real_citations_are_kept_separate(self):
        cases = {
            'Welchov t(173,59) = 1,77': [], 'F(4; 44,89) = 3,57': [],
            'korelacija r(58) = 0,31': [], 'χ(2) = 4,5': [], 'razlika (44,08)': [],
            'stigma (12,40)': [12, 40], 'korelacija iznosi 0,53 (21).': [21],
            'znanje(12,40)': [12, 40], 'Dokaz (99).': [99], 'Dokaz (1–5).': [1, 2, 3, 4, 5],
        }
        for text, expected in cases.items():
            with self.subTest(text=text):
                actual = [n for _pos, group in COMMON.find_vancouver_citations(text) for n in group]
                self.assertEqual(actual, expected)

    def test_plural_reference_preserves_unmentioned_table_control(self):
        import check_uputnice as module
        path = make_docx(self.dir / 'references.docx', [
            'Rezultati u Tablicama 7. i 8. pokazuju razliku.',
            'Tablica 5. Naslov', 'Tablica 7. Naslov', 'Tablica 8. Naslov'])
        data = module.analiziraj(path)
        self.assertEqual(data['nespomenuti'], [('Tablica', [5])])
        self.assertEqual(data['u_prazno'], [])

    def test_plural_reference_missing_target_still_fails(self):
        import check_uputnice as module
        path = make_docx(self.dir / 'absent_table.docx', ['Vidi Tablice 1 i 99.', 'Tablica 1. Naslov'])
        data = module.analiziraj(path)
        self.assertEqual(data['u_prazno'][0][1], [99])
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(module.main(path), 1)

    def test_hypothesis_verdict_synonyms_do_not_invent_a_verdict(self):
        import check_hipoteze as module
        for verdict in ['podržana', 'nije podržana', 'poduprta', 'potkrijepljena', 'prihvaćena']:
            with self.subTest(verdict=verdict):
                self.assertIsNotNone(module.PRESUDA.search('Hipoteza H1 je ' + verdict + '.'))
        self.assertIsNone(module.PRESUDA.search('Hipoteza H4 odnosi se na radni staž.'))
        self.assertIsNone(module.PRESUDA.search('Rezultati podržavaju raspravu o H4.'))

    def test_alpha_declaration_forms_and_negative_controls(self):
        import check_statistika as module
        for text, expected in [
            ('Razina statističke značajnosti postavljena je na p < 0,05.', .05),
            ('Korištena je α = 0,01.', .01), ('Razina značajnosti iznosi 0,05.', .05),
            ('Statistički značajnim smatra se p < 0,01.', .01),
            ('Rezultati su prikazani u tablicama.', None),
            ('Razlika nije bila značajna (p = 0,190).', None),
        ]:
            with self.subTest(text=text):
                self.assertEqual(module._deklaracija_praga(text)[0], expected)

    def test_category_sums_do_not_split_decimals_or_ranges(self):
        import numbers_inventory as module
        self.assertEqual(module._pribrojnici('129 ispitanika (63,5 %).'), [129])
        self.assertEqual(module._pribrojnici('između 8 % i 11,2 % izdataka'), [])
        text = 'Ukupno 100 ispitanika: 40 muškaraca, 35 žena i 25 ostalih.'
        self.assertEqual(module._pribrojnici(text), [100, 40, 35, 25])
        self.assertTrue(module.zbroj_kategorija([text]))
        self.assertEqual(module.zbroj_kategorija([
            '129 ispitanika (63,5 %) ima od 6 do 10 godina staža (50,7 %).']), [])

    def test_reference_title_quotes_do_not_disable_body_checks(self):
        import check_typography as module
        path = make_docx(self.dir / 'quotes.docx', ['Program „Neovisno življenje” je pokrenut.'],
                         ['1. Autor. The “do not resuscitate” order. 2020.'])
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            module.main(path)
        self.assertNotIn('dva oblika zatvarajućeg', buf.getvalue())
        self.assertNotIn('nesparen', buf.getvalue())
        bad = make_docx(self.dir / 'bad_quotes.docx', ['Program "Neovisno življenje" je pokrenut.'])
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            module.main(bad)
        self.assertIn('⚠ ravni navodnici', buf.getvalue())

    def test_inventory_missing_document_is_unmeasured(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            code = INV.main([str(self.dir / 'missing.docx')])
        self.assertEqual(code, 3, buf.getvalue())

    def test_inventory_missing_bibliography_is_unmeasured(self):
        result, error = INV.analiza(self.non_hks_doc())
        self.assertIsNone(result)
        self.assertTrue(error)

    def test_inventory_zero_batch_is_explicitly_rejected(self):
        path = make_docx(self.dir / 'batch.docx', ['Uzorak ima 42 ispitanika (1).'], [REFS[0]])
        with contextlib.redirect_stdout(io.StringIO()):
            code = INV.main([path, '--dosjei', str(self.dir / 'dossiers'), '--po-seriji', '0'])
        self.assertEqual(code, 2)
        self.assertFalse((self.dir / 'dossiers').exists())

    def test_inventory_dossiers_contain_actual_numeric_and_descriptive_claims(self):
        path = make_docx(self.dir / 'claims.docx', [
            'Uzorak ima 42 ispitanika (1).', 'Model opisuje zajednicu (2).'], REFS[:2])
        output = self.dir / 'claims.json'
        dossiers = self.dir / 'dossiers'
        with contextlib.redirect_stdout(io.StringIO()):
            code = INV.main([path, '--json', str(output), '--dosjei', str(dossiers), '--po-seriji', '1'])
        self.assertEqual(code, 0)
        data = json.loads(output.read_text(encoding='utf-8'))
        self.assertTrue(data['1']['brojcane'])
        self.assertTrue(data['2']['opisne'])
        files = list(dossiers.glob('*.md'))
        self.assertEqual(len(files), 1)
        self.assertIn('42 ispitanika', files[0].read_text(encoding='utf-8'))

    def test_manual_discussion_review_is_unmeasured(self):
        path = self.dir / 'discussion.docx'
        doc = Document()
        doc.add_heading('1. UVOD', level=1)
        doc.add_paragraph('Uvodni tekst.')
        doc.add_heading('2. RASPRAVA', level=1)
        doc.add_paragraph('Ovaj odlomak tek treba sadržajno provjeriti.')
        doc.save(path)
        result = HKS.analiza(str(path))
        finding = next(x for x in result['nalazi'] if 'prvi odlomak sažima rezultate?' in x['poruka'])
        self.assertIsNone(finding['ok'])
        output = io.StringIO()
        with contextlib.redirect_stdout(output):
            HKS.ispisi(result)
        line = next(x for x in output.getvalue().splitlines() if 'prvi odlomak sažima' in x)
        self.assertNotIn('✅', line)

    def test_inventory_unsupported_style_is_not_forced_to_vancouver(self):
        for text in ['Model je opisan (Horvat, 2020).',
                     'Model je opisan [1] i potvrđen (1).',
                     'Model je opisan bez numeričkog citata.']:
            with self.subTest(text=text):
                path = make_docx(self.dir / 'unsupported.docx', [text], [REFS[0]])
                result, error = INV.analiza(path)
                self.assertIsNone(result)
                self.assertIn('stil', error.lower())
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(INV.main([path]), 3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
